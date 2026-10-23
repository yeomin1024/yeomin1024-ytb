# -*- coding: utf-8 -*-
# VERSION: v1.1 — 2026-10-23 — comment_clusters.py v1.1: 나라별 표·몫 · 48주(96 기간) · 번역 장부(merge·todo) · 영상 순위 두 가지(하루 조회수·조회수).
# v1.0 — 2026-10-23 — tools/comment_clusters.py v1.0 시험: 기간(한 주씩·누적·직접 지정)·묶음 통계(댓글 수·좋아요·몫)·영상 순서(하루 조회수)·
#          묶기 씨앗 고정·엑셀 시트 (임시 폴더만 — 저장소 결과·네트워크·GitHub 안 건드림).
"""python3 -m unittest test_comment_clusters  (tools/tests 에서)"""
import os
import random
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
for k in ("GITHUB_TOKEN", "GH_TOKEN"):
    os.environ.pop(k, None)
import comment_clusters as cc  # noqa: E402


class Periods(unittest.TestCase):
    def test_default(self):
        names = [p[0] for p in cc.periods(3)]
        self.assertEqual(names, ["1주 전", "2주 전", "3주 전", "1~2주 전", "1~3주 전", "전체"])
        ps = cc.periods(cc.DEFAULT_WEEKS)
        self.assertEqual(len(ps), 96)                               # 48 한 주 + 47 누적 + 전체
        self.assertEqual((ps[47][0], ps[-2][0]), ("48주 전", "1~48주 전"))

    def test_custom(self):
        ps = cc.parse_periods("1, 2-4 ,1~3", 4)
        self.assertEqual(ps, [("1주 전", 1, 1), ("2~4주 전", 2, 4), ("1~3주 전", 1, 3), ("전체", None, None)])
        self.assertEqual(cc.parse_periods("", 2), cc.periods(2))
        with self.assertRaises(SystemExit):
            cc.parse_periods("3-1", 4)
        with self.assertRaises(SystemExit):
            cc.parse_periods("지난주", 4)

    def test_week_boundaries(self):
        self.assertEqual(cc.week_of_age(0), 1)
        self.assertEqual(cc.week_of_age(6.99), 1)                 # 0~6일 전 = 1주 전
        self.assertEqual(cc.week_of_age(7), 2)
        self.assertIsNone(cc.week_of_age(None))
        self.assertIsNone(cc.week_of_age(-1))                     # 미래 날짜 — 어느 주에도 넣지 않음
        p13 = ("1~3주 전", 1, 3)
        self.assertTrue(cc.in_period(3, p13))
        self.assertFalse(cc.in_period(4, p13))
        self.assertFalse(cc.in_period(None, p13))
        self.assertTrue(cc.in_period(None, ("전체", None, None)))


class Tokens(unittest.TestCase):
    def test_en(self):
        self.assertEqual(cc.en_tokens(["The dividends are great, buy SCHD https://x.y @bob"])[0], ["dividend", "great", "buy", "schd"])

    def test_ko(self):
        try:
            import kiwipiepy  # noqa: F401
        except ImportError:
            self.skipTest("kiwipiepy 없음")
        w = cc.ko_tokens(["국민연금 손실 너무 크네요 ㅋㅋ", "ㅋㅋㅋㅋ👍"])
        self.assertIn("국민연금", w[0])
        self.assertIn("크다", w[0])                                # 형용사 원형
        self.assertNotIn("너무", w[0])                             # 멈춤 낱말
        self.assertEqual(w[1], [])                                 # 내용 낱말 없음 → 0번 묶음

    def test_clean_zero_width(self):
        self.assertEqual(cc.clean("a​  b\n c"), "a b c")


def C(vid, cl, likes, text="댓글 내용 여덟 글자 이상", lang="한국어", sim=0.5, week=1, country=None):
    return {"vid": vid, "cluster": cl, "likes": likes, "text": text, "lang_group": lang, "sim": sim, "week": week,
            "country": country or ("한국" if lang == "한국어" else "미국")}


class FakeTr(cc.Translations):
    def __init__(self, known=None):
        self.ko, self.en, self.missing, self.used = dict(known or {}), {}, {}, set()


class Stats(unittest.TestCase):
    INFO = {"한국어": {0: {"label": cc.EMPTY}, 1: {"label": "연금 · 손실"}, 2: {"label": "배당"}}, "영어": {1: {"label": "dividend"}}}

    def test_counts_likes_shares(self):
        sel = [C("a", 1, 10, sim=0.9, text="가장 가운데 댓글입니다"), C("a", 1, 30), C("b", 2, 60), C("b", 0, 0, text="ㅋㅋ"),
               C("c", 1, 5, lang="영어", text="dividend comment here")]
        tr = FakeTr({cc.Translations.key("dividend comment here"): "배당 댓글"})
        rows = {(r["나라"], r["묶음"]): r for r in cc.cluster_table(cc.frame(sel), sel, self.INFO, tr)}
        r = rows[("한국", "KO-001")]
        self.assertEqual(r["댓글 수"], 2)
        self.assertAlmostEqual(r["댓글 몫"], 0.5)                  # 한국 4개 중 2
        self.assertEqual(r["좋아요 합"], 40)
        self.assertAlmostEqual(r["좋아요 몫"], 0.4)                # 한국 좋아요 100 중 40
        self.assertEqual(r["영상 수"], 1)
        self.assertEqual(r["대표 댓글 1"], "가장 가운데 댓글입니다")   # 한가운데(sim 큰 것)
        self.assertTrue(r["좋아요 1위 댓글"].startswith("👍10 "))   # 👍30 은 한가운데에서 먼 절반(sim 0.5 < 중앙값 0.7) → 보기에서 뺌, 좋아요 합에는 듦
        self.assertNotIn("대표 댓글 1 번역", r)                     # 한국 시트엔 번역 칸 없음
        self.assertEqual(rows[("한국", "KO-000")]["대표 댓글 1"], "")  # 0번 묶음은 대표 없음 (칸은 있음)
        us = rows[("미국", "EN-001")]
        self.assertAlmostEqual(us["댓글 몫"], 1.0)                  # 몫은 나라마다 따로
        self.assertEqual(us["대표 댓글 1 번역"], "배당 댓글")        # 장부에 있는 번역
        self.assertEqual(us["주제 낱말 번역"], "")                  # 장부에 없음 → 비고 할 일로
        self.assertIn(cc.Translations.key("dividend"), tr.missing)
        order = [(r["나라"], r["묶음"]) for r in cc.cluster_table(cc.frame(sel), sel, self.INFO, tr)]
        self.assertEqual(order[0], ("한국", "KO-001"))             # 한국 먼저, 댓글 많은 순
        self.assertEqual(order[-1], ("미국", "EN-001"))

    def test_period_mask(self):
        sel = [C("a", 1, 1, week=1), C("a", 1, 1, week=3), C("a", 1, 1, week=None)]
        df = cc.frame(sel)
        self.assertEqual(cc.period_mask(df, ("1~2주 전", 1, 2)).tolist(), [True, False, False])
        self.assertEqual(cc.period_mask(df, ("전체", None, None)).tolist(), [True, True, True])   # 날짜 모름도 전체엔 듦
        self.assertEqual([cc.period_kind(p) for p in (("x", 1, 1), ("x", 1, 3), ("x", 2, 4), ("x", None, None))], ["한 주", "누적", "범위", "전체"])

    def test_video_table_ranks(self):
        vids = [{"id": i, "title": i, "channel": "", "subs": 1, "views": views, "vpd": vpd, "published": None, "age_days": age,
                 "lang_group": "한국어", "country": "한국", "group": "", "minutes": 5, "channel_country": "KR"}
                for i, vpd, views, age in (("a", 10.0, 900, 3), ("b", 500.0, 100, 10), ("c", None, 50, 400))]
        vids.append({"id": "e", "title": "Buy the dip", "channel": "", "subs": 1, "views": 10, "vpd": 1.0, "published": None, "age_days": 1,
                     "lang_group": "영어", "country": "미국", "group": "", "minutes": 5, "channel_country": "US"})
        cms = [C("a", 1, 3), C("b", 2, 7), C("b", 2, 1, week=2), C("b", 1, 0, week=1), C("b", 0, 0, text="ㅋㅋ"), C("b", 0, 0, text="👍"),
               C("b", 0, 0, text="ㅎㅎ"), C("e", 1, 2, lang="영어", text="great dip")]
        tr = FakeTr()
        rows, cross = cc.video_table(vids, cms, self.INFO, 48, tr, by="comment")
        r = {x["제목"]: x for x in rows}
        self.assertEqual((r["b"]["순위(하루 조회수)"], r["a"]["순위(하루 조회수)"], r["c"]["순위(하루 조회수)"]), (1, 2, 3))
        self.assertEqual((r["a"]["순위(조회수)"], r["b"]["순위(조회수)"], r["c"]["순위(조회수)"]), (1, 2, 3))   # 조회수 순은 따로
        self.assertEqual(r["Buy the dip"]["순위(조회수)"], 1)        # 순위는 나라 안에서
        self.assertEqual(r["b"]["주"], "2주 전")
        self.assertEqual(r["c"]["주"], "그 전")                     # 48주보다 오래됨
        self.assertEqual(r["b"]["시청자 댓글 수"], 6)
        self.assertEqual(r["b"]["낱말 없음(이모지·ㅋㅋ) 댓글 수"], 3)
        self.assertTrue(r["b"]["많은 묶음 1"].startswith("KO-002 배당 — 2개"))  # 0번(3개)이 더 많아도 내용 있는 묶음만
        self.assertEqual((r["b"]["1주 전 댓글 수"], r["b"]["2주 전 댓글 수"], r["b"]["2주 전 좋아요 합"], r["b"]["48주 전 댓글 수"]), (5, 1, 1, 0))
        self.assertNotIn("제목 번역", r["a"])
        self.assertIn("제목 번역", r["Buy the dip"])
        self.assertEqual(set(tr.missing.values()), {"Buy the dip", "great dip"})
        b2 = next(x for x in cross if x["제목"] == "b" and x["묶음"] == "KO-002")
        self.assertAlmostEqual(b2["그 영상 좋아요 중 몫"], 1.0)
        self.assertEqual(b2["순위(하루 조회수)"], 1)
        self.assertNotIn("1주 전 댓글 수", cc.video_table(vids, cms, self.INFO, 48, FakeTr())[0][0])   # 영상 올린 날 기준이면 주별 칸 없음


class TranslationBook(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def test_roundtrip_merge_todo(self):
        book = cc.Translations(self.tmp / "tr.tsv")
        self.assertEqual(book.get("Buy  the\ndip"), "")            # 정리한 글로 key
        self.assertEqual(book.key("Buy  the\ndip"), book.key("Buy the dip"))
        self.assertEqual(book.write_todo(self.tmp / "todo.tsv"), 1)
        todo = (self.tmp / "todo.tsv").read_text("utf-8").splitlines()
        self.assertEqual(todo[0], "key\ten\tko")
        k = todo[1].split("\t")[0]
        src = self.tmp / "done.tsv"
        src.write_text(f"key\ten\tko\n{k}\tBuy the dip\t저점 매수\nbadkey\tBuy the dip\t틀린 key\n{k[:-1]}x\t\t\n", "utf-8")
        added, skipped = book.merge(src)
        self.assertEqual((added, len(skipped)), (1, 2))             # key 와 글이 안 맞음 · 번역 없음 → 건너뜀
        book.save()
        again = cc.Translations(self.tmp / "tr.tsv")
        self.assertEqual(again.get("Buy the dip"), "저점 매수")
        self.assertEqual(again.write_todo(self.tmp / "todo.tsv"), 0)
        self.assertFalse((self.tmp / "todo.tsv").exists())          # 할 일이 없으면 파일을 지움

    def test_quotes_kept(self):
        book = cc.Translations(self.tmp / "tr.tsv")
        k = book.key('"Buy" the dip')
        (self.tmp / "q.tsv").write_text(f'key\tko\n{k}\t"저점" 매수 "꼭"\n', "utf-8")    # 따옴표로 시작·끝 — csv 모듈이면 잘못 읽힘
        self.assertEqual(book.merge(self.tmp / "q.tsv"), (1, []))
        book.save()
        self.assertEqual(cc.Translations(self.tmp / "tr.tsv").get('"Buy" the dip'), '"저점" 매수 "꼭"')


class ClusterSeed(unittest.TestCase):
    def test_same_seed_same_labels(self):
        rng = random.Random(1)
        topics = (["dividend", "income", "yield", "schd"], ["crash", "bubble", "recession", "panic"], ["bitcoin", "crypto", "coin", "halving"])
        rows = [{"text": " ".join(rng.choice(t) for _ in range(4)) + " stock"} for t in topics for _ in range(60)] + [{"text": "😂😂"}] * 5
        a = cc.cluster_lang("영어", rows, 3, 0.5)
        b = cc.cluster_lang("영어", rows, 3, 0.5)
        self.assertEqual(a[0].tolist(), b[0].tolist())
        self.assertTrue(all(x == 0 for x in a[0][-5:]))           # 낱말 없는 댓글 = 0번
        groups = [set(a[0][i * 60:(i + 1) * 60].tolist()) for i in range(3)]
        self.assertTrue(all(len(g) == 1 for g in groups))          # 같은 주제 댓글은 한 묶음
        self.assertEqual(len(set.union(*groups)), 3)
        self.assertTrue(a[2][1]["terms"])                          # 주제 낱말이 있음


class Xlsx(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def test_sheets(self):
        import openpyxl
        p = self.tmp / "x.xlsx"
        cc.write_xlsx(p, [("한국 기간별", [{"묶음": "KO-001", "댓글 몫": 0.25, "댓글 수": 1200}]),
                          ("그 밖 나라 영상별(하루 조회수 순)", [{"제목": "t", "링크": "https://www.youtube.com/watch?v=abc"}])], ["안내 한 줄"])
        wb = openpyxl.load_workbook(p)
        self.assertEqual(wb.sheetnames, ["안내", "한국 기간별", "그 밖 나라 영상별(하루 조회수 순)"])
        ws = wb["한국 기간별"]
        self.assertEqual(ws["B2"].number_format, "0.0%")
        self.assertEqual(ws["C2"].number_format, "#,##0")
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(wb["그 밖 나라 영상별(하루 조회수 순)"]["B2"].hyperlink.target, "https://www.youtube.com/watch?v=abc")


if __name__ == "__main__":
    unittest.main()
