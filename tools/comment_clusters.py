#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# VERSION: v1.0 — 2026-10-23 — 새로: 내용이 비슷한 댓글끼리 묶어 기간별·영상별 통계를 엑셀로. 사용자 지시 2026-10-23 "아니 분류를 댓글 내용상 비슷한것
#          끼리 카운트를 세고 그 대표 댓글 내용을 알려줘야 해 각 1주전,2주전, 3주전... 이랑 1주전~2주전, 1주전~3주전 이런식으로 기간을 조정해서 통계 낼
#          수 있도록 하고 댓글 내용 카운트, 좋아요 수를 각각 통계를 내야하고 업로드 기간 대비 조회수가 가장 높은 영상 순으로 정렬해서 그 영상별 댓글 관련
#          통계도 내야해 그걸 엑셀 파일로 정리하도록 코드 수정해봐".
#          · 묶기: 언어 묶음(한국어·영어·그 밖)마다 내용 낱말(한국어 = Kiwi 형태소의 명사·동사·형용사 원형, 영어 = 낱말, 그 밖 = 글자 조각) → TF-IDF →
#            SVD → MiniBatchKMeans (씨앗 고정) → 너무 큰 묶음(--split 비율 넘음)은 다시 나눔. 내용 낱말이 없는 댓글(이모지·ㅋㅋ)은 따로 한 묶음.
#            '비슷하다' = 같은 낱말을 쓴다는 뜻 — 뜻이 같아도 다른 낱말이면 다른 묶음일 수 있다 (결과 '안내' 시트).
#          · 기간: 영상 올린 날(기본) 또는 댓글 단 날(--by comment) 기준 — 한 주씩(1주 전·2주 전…) · 누적(1~2주 전·1~3주 전…) · 전체. --weeks 로 몇 주까지,
#            --periods "1,2,1-3,2-4" 로 직접.
#          · 통계: 묶음마다 댓글 수·몫, 좋아요 합·몫·평균, 영상 수, 대표 댓글(묶음 한가운데에 가까운 것)·좋아요 많은 댓글.
#          · 영상별: 하루 조회수(받은 때 조회수 ÷ 올린 지 날수) 큰 순 — 댓글 수·좋아요·많은 묶음 3개·좋아요 1위 댓글 + 영상×묶음 표.
#          할당량 0. ※ 연구·교육용 — 투자 권유가 아님.
"""내용이 비슷한 댓글 묶음 — 기간별·영상별 통계 엑셀

  python3 tools/comment_clusters.py stock                      # stock/result/outliers/comment_clusters.xlsx (+ .md 요약)
  python3 tools/comment_clusters.py stock --weeks 6            # 1~6주 전 한 주씩 + 1~2…1~6주 전 누적
  python3 tools/comment_clusters.py stock --periods "1,2,1-3,2-4"   # 기간 직접 (2-4 = 2~4주 전)
  python3 tools/comment_clusters.py stock --by comment         # 기간 = 댓글 단 날 기준 (기본 = 영상 올린 날)
  python3 tools/comment_clusters.py stock --k 한국어=60,영어=40,그밖=15 --split 0.03
"""
import argparse
import json
import logging
import math
import re
import sys
import time
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parent))
import comment_types as ct            # noqa: E402  lang_group (영상 언어 묶음) — 같은 규칙
import title_comment_link as tcl      # noqa: E402  country_of (한국·미국) · 영어 STOP_EN

VERSION = "v1.0"
ROOT = Path(__file__).resolve().parent.parent
LOG = logging.getLogger("cclu")
SEED = 20261023
LANGS = ("한국어", "영어", "그 밖 언어")
DEFAULT_K = {"한국어": 60, "영어": 40, "그 밖 언어": 15}
SPLIT_SHARE = 0.03            # 한 묶음이 그 언어 댓글의 이 몫을 넘으면 다시 나눔
SVD_DIM = 100
MIN_DF = 5
REP_TYPICAL = 2               # 대표 댓글 (묶음 한가운데에 가까운 것)
REP_LIKED = 3                 # 좋아요 많은 댓글
TEXT_MAX = 300                # 엑셀 칸에 넣는 댓글 글자 수
EMPTY = "내용 낱말 없음 (이모지·ㅋㅋ·짧은 반응)"
KO_KEEP = {"NNG", "NNP", "SL", "VV", "VA", "XR", "SH"}
KO_STOP = set("것 수 거 등 때 저 이 그 좀 더 안 못 잘 하다 되다 있다 없다 같다 보다 이다 아니다 주다 가다 오다 알다 말 사람 우리 지금 진짜 정말 "
              "그냥 너무 많이 이제 왜 뭐 거기 여기 저기 이거 그거 저거 분 님 씨 들 중 번 개 년 월 일 원 만 억 조".split())
EN_STOP = tcl.STOP_EN | set("like just get got one would could also really much many even still people think know going want lol im dont "
                            "thats youre theyre cant wont didnt doesnt isnt arent wasnt sir bro guys".split())
PREFIX = {"한국어": "KO", "영어": "EN", "그 밖 언어": "OT"}
KO_ONE = set("돈 집 빚 금 땅 욕 법 꿈 삶 탓 꼴 짓 죄 운 표 장 빵 밥 술 약 벌".split())   # 한 글자 명사는 이것만 남김 (드·콜·린·주·절 같은 조각은 뺌)
ZW = re.compile(r"[​-‏⁠﻿]")


def log(stage, level=logging.INFO, **kv):
    LOG.log(level, "[%s] [%s] %s", stage, datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S"),
            " ".join(f"{k}={json.dumps(v, ensure_ascii=False) if isinstance(v, str) and (' ' in v or not v) else v}" for k, v in kv.items()))


def clean(t, n=TEXT_MAX):
    return re.sub(r"\s+", " ", ZW.sub("", str(t or ""))).strip()[:n]


def parse_dt(s):
    try:
        d = datetime.fromisoformat(str(s).replace("Z", "+00:00"))
        return d if d.tzinfo else d.replace(tzinfo=timezone.utc)
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------- 자료
def load(topic, now):
    """승인한 영상 → (영상 목록, 시청자 댓글 목록). 하루 조회수 = 영상 정보를 받은 때 조회수 ÷ 올린 날부터 그때까지 날수 (최소 1일)."""
    out = ROOT / topic / "result" / "outliers"
    an = json.loads((out / "analysis.json").read_text("utf-8")) if (out / "analysis.json").is_file() else {}
    success_ids = {c["id"] for c in an.get("success_cases") or []}
    vpd_ids = {c["id"] for c in (an.get("vpd_top") or {}).get("cases") or []}
    chp = tcl.channels_path(out)
    ch = json.loads(chp.read_text("utf-8")) if chp.is_file() else {}
    videos, comments = [], []
    for d, meta, info, cms in ct.iter_videos(out):
        vid = meta.get("id")
        pub = parse_dt(info.get("published") or meta.get("published_at") or meta.get("published"))
        fetched = parse_dt(info.get("fetched")) or now
        views = int(info.get("view_count") or meta.get("views") or 0)
        days = max((fetched - pub).total_seconds() / 86400, 1.0) if pub else None
        v = {"id": vid, "title": meta.get("title") or info.get("title") or "", "channel": meta.get("channel") or "", "subs": meta.get("subs"),
             "views": views, "vpd": round(views / days, 1) if days else None, "published": pub, "age_days": (now - pub).total_seconds() / 86400 if pub else None,
             "lang_group": ct.lang_group(meta, info),
             "country": tcl.country_of({"lang": info.get("lang") or meta.get("video_lang"), "audio_lang": info.get("audio_lang"), "channel_id": meta.get("channel_id")}, ch) or "그 밖 나라",
             "group": "성공 사례" if vid in success_ids else ("하루 조회수 상위" if vid in vpd_ids else "그 밖의 승인 영상"),
             "minutes": meta.get("minutes")}
        videos.append(v)
        for m in cms:
            if m["creator"] or not ZW.sub("", m["text"]).strip():
                continue
            comments.append({"vid": vid, "cid": m["cid"], "text": m["text"], "likes": m["likes"], "reply": m["reply"], "lang_group": v["lang_group"],
                             "published": parse_dt(m.get("published"))})            # 댓글 단 때 (--by comment 일 때만 씀)
    log("LOAD", topic=topic, videos=len(videos), comments=len(comments), by_lang=dict(Counter(c["lang_group"] for c in comments)))
    return out, videos, comments


# ---------------------------------------------------------------- 낱말
_KIWI = None


def ko_tokens(texts):
    global _KIWI
    from kiwipiepy import Kiwi
    if _KIWI is None:
        _KIWI = Kiwi()
    res = []
    for toks in _KIWI.tokenize([clean(re.sub(r"https?://\S+|@\S+", " ", t), 500) for t in texts]):
        words = []
        for t in toks:
            if t.tag not in KO_KEEP:
                continue
            w = t.form + ("다" if t.tag in ("VV", "VA") else "")
            if w in KO_STOP or (t.tag == "SL" and len(w) < 2) or (len(w) == 1 and w not in KO_ONE):
                continue
            words.append(w.lower() if t.tag == "SL" else w)
        res.append(words)
    return res


def en_tokens(texts):
    out = []
    for t in texts:
        ws = re.findall(r"[a-z][a-z0-9&']+", clean(re.sub(r"https?://\S+|@\S+", " ", t), 500).lower())
        out.append([w[:-1] if w.endswith("s") and len(w) > 3 and not w.endswith("ss") else w for w in (x.replace("'", "") for x in ws)
                    if w not in EN_STOP and len(w) >= 3])
    return out


def other_tokens(texts):
    """그 밖 언어 (힌디 로마자·일본어·독일어 …): 낱말 + 띄어쓰기 없는 글(일본어·중국어)은 두 글자 조각."""
    out = []
    for t in texts:
        s = clean(re.sub(r"https?://\S+|@\S+", " ", t), 500).lower()
        ws = [w for w in re.findall(r"\w+", s) if len(w) >= 3 and not w.isdigit() and w not in EN_STOP]
        cjk = re.sub(r"[^぀-ヿ一-鿿]", "", s)
        ws += [cjk[i:i + 2] for i in range(len(cjk) - 1)]
        out.append(ws)
    return out


TOKENIZERS = {"한국어": ko_tokens, "영어": en_tokens, "그 밖 언어": other_tokens}


# ---------------------------------------------------------------- 묶기
def cluster_lang(lang, rows, k, split_share, seed=SEED):
    """한 언어 묶음의 댓글 → 묶음 번호(0 = 내용 낱말 없음)·주제 낱말·대표 정도(한가운데와의 닮음). 씨앗 고정 — 다시 돌려도 같음."""
    from sklearn.cluster import MiniBatchKMeans
    from sklearn.decomposition import TruncatedSVD
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.preprocessing import normalize
    t0 = time.time()
    toks = TOKENIZERS[lang]([r["text"] for r in rows])
    has = np.array([len(x) > 0 for x in toks])
    idx = np.where(has)[0]
    labels = np.zeros(len(rows), dtype=int)            # 0 = 내용 낱말 없음
    sim = np.zeros(len(rows))
    info = {0: {"terms": [], "label": EMPTY}}
    if len(idx) < max(k * 3, 30):
        log("CLUSTER", logging.WARNING, lang=lang, comments=len(rows), with_words=len(idx), note="댓글이 적어 묶지 않음 — 모두 한 묶음")
        for i in idx:
            labels[i] = 1
        info[1] = {"terms": [], "label": "(댓글이 적어 묶지 않음)"}
        return labels, sim, info
    vec = TfidfVectorizer(tokenizer=str.split, preprocessor=None, lowercase=False, token_pattern=None,
                          min_df=min(MIN_DF, max(1, len(idx) // 1000)), max_df=0.3, sublinear_tf=True)
    X = vec.fit_transform([" ".join(toks[i]) for i in idx])
    nonzero = np.asarray(X.getnnz(axis=1)).ravel() > 0          # min_df·max_df 로 낱말이 다 빠진 댓글 → 낱말 없음 묶음
    idx, X = idx[nonzero], X[nonzero]
    dim = min(SVD_DIM, X.shape[1] - 1)
    Z = normalize(TruncatedSVD(n_components=dim, random_state=seed).fit_transform(X))
    k = min(k, max(2, len(idx) // 50))
    km = MiniBatchKMeans(n_clusters=k, random_state=seed, batch_size=4096, n_init=3).fit(Z)
    lab = km.labels_.copy()
    cent = [c for c in km.cluster_centers_]
    limit = max(int(split_share * len(lab)), 50)
    for _round in range(5):                                     # 너무 큰 묶음은 다시 나눔 (최대 5번 돌기)
        changed = False
        for c in range(len(cent)):
            m = np.where(lab == c)[0]
            if len(m) <= limit:
                continue
            k2 = int(math.ceil(len(m) / (limit * 0.5)))
            sub = MiniBatchKMeans(n_clusters=k2, random_state=seed, batch_size=4096, n_init=3).fit(Z[m])
            if np.bincount(sub.labels_, minlength=k2).max() > 0.9 * len(m):
                continue                                         # 더 나뉘지 않음 — 그대로 둠
            for j in range(k2):
                mj = m[sub.labels_ == j]
                if j == 0:
                    cent[c] = sub.cluster_centers_[0]
                else:
                    lab[mj] = len(cent)
                    cent.append(sub.cluster_centers_[j])
            changed = True
        if not changed:
            break
    cent = np.array([v / (np.linalg.norm(v) or 1.0) for v in cent])
    terms = np.array(vec.get_feature_names_out())
    order = np.argsort(-np.bincount(lab, minlength=len(cent)))  # 큰 묶음부터 1, 2, 3 …
    remap = {int(c): n + 1 for n, c in enumerate(order)}
    base = np.asarray(X.mean(axis=0)).ravel()
    for c in range(len(cent)):
        m = lab == c
        if not m.any():
            continue
        mean = np.asarray(X[m].mean(axis=0)).ravel()
        lift = mean - base                                      # 그 묶음에서 유난히 많은 낱말 (어디에나 많은 '주식' 같은 낱말은 뒤로)
        top = [str(terms[j]) for j in np.argsort(-lift)[:6] if lift[j] > 0]
        info[remap[c]] = {"terms": top, "label": " · ".join(top[:5])}
    labels[idx] = [remap[int(c)] for c in lab]
    sim[idx] = np.einsum("ij,ij->i", Z, cent[lab])
    log("CLUSTER", lang=lang, comments=len(rows), with_words=len(idx), no_words=len(rows) - len(idx), vocab=X.shape[1], k_start=k,
        clusters=len(info) - 1, largest=int(np.bincount(lab).max()), split_limit=limit, seed=seed, seconds=round(time.time() - t0, 1))
    return labels, sim, info


# ---------------------------------------------------------------- 기간
def periods(weeks):
    """(이름, 시작 주, 끝 주) — 한 주씩 1…N, 누적 1~2…1~N, 전체(None)."""
    ps = [(f"{w}주 전", w, w) for w in range(1, weeks + 1)]
    ps += [(f"1~{w}주 전", 1, w) for w in range(2, weeks + 1)]
    ps.append(("전체", None, None))
    return ps


def parse_periods(spec, weeks):
    """--periods "1,2,1-3,2-4" → [("1주 전",1,1), ("2주 전",2,2), ("1~3주 전",1,3), ("2~4주 전",2,4), ("전체",None,None)]. 비면 periods(weeks)."""
    if not (spec or "").strip():
        return periods(weeks)
    ps = []
    for part in spec.split(","):
        part = part.strip().replace("~", "-")
        if not part:
            continue
        m = re.fullmatch(r"(\d+)(?:-(\d+))?", part)
        if not m:
            raise SystemExit(f"[PERIOD] ❌ 기간 {part!r} 을 모름 — 예: 1,2,1-3,2-4 (숫자 = 몇 주 전)")
        a, b = int(m.group(1)), int(m.group(2) or m.group(1))
        if not 1 <= a <= b:
            raise SystemExit(f"[PERIOD] ❌ 기간 {part!r}: 1 이상, 앞 ≤ 뒤")
        ps.append((f"{a}주 전" if a == b else f"{a}~{b}주 전", a, b))
    ps.append(("전체", None, None))
    return ps


def week_of_age(age_days):
    if age_days is None or age_days < 0:
        return None
    return int(age_days // 7) + 1


def in_period(week, p):
    _, a, b = p
    return True if a is None else (week is not None and a <= week <= b)


# ---------------------------------------------------------------- 통계
def cluster_id(lang, cl):
    return f"{PREFIX.get(lang, 'OT')}-{cl:03d}"


def stats_rows(sel, info_by_lang, typical_n=REP_TYPICAL, liked_n=REP_LIKED):
    """고른 댓글 → 묶음별 줄 (언어마다 몫을 따로 셈)."""
    by = defaultdict(list)
    for c in sel:
        by[(c["lang_group"], c["cluster"])].append(c)
    tot_n = Counter(c["lang_group"] for c in sel)
    tot_l = Counter()
    for c in sel:
        tot_l[c["lang_group"]] += c["likes"]
    out = []
    for (lg, cl), cs in by.items():
        likes = sum(c["likes"] for c in cs)
        typical = [clean(c["text"]) for c in sorted(cs, key=lambda c: -c["sim"]) if len(clean(c["text"])) >= 8][:typical_n] if cl else []
        med = float(np.median([c["sim"] for c in cs])) if cl else 0.0
        fit = [c for c in cs if c["sim"] >= med] if cl else cs  # 묶음에 잘 맞는 절반 (낱말 하나만 겹친 댓글은 보기에서 뺌)
        liked = [f"👍{c['likes']:,} {clean(c['text'])}" for c in sorted(fit, key=lambda c: -c["likes"])[:liked_n]]
        out.append({"언어": lg, "묶음": cluster_id(lg, cl), "주제 낱말": info_by_lang[lg].get(cl, {}).get("label", ""),
                    "댓글 수": len(cs), "댓글 몫": len(cs) / tot_n[lg] if tot_n[lg] else None, "좋아요 합": likes,
                    "좋아요 몫": likes / tot_l[lg] if tot_l[lg] else None, "평균 좋아요": round(likes / len(cs), 2), "영상 수": len({c["vid"] for c in cs}),
                    **{f"대표 댓글 {i + 1}": typical[i] if i < len(typical) else "" for i in range(typical_n)},       # 칸은 늘 만듦 (열 순서 고정)
                    **{f"좋아요 {i + 1}위 댓글": liked[i] if i < len(liked) else "" for i in range(liked_n)}})
    lang_order = {lg: n for n, lg in enumerate(LANGS)}
    return sorted(out, key=lambda r: (lang_order.get(r["언어"], 9), -r["댓글 수"]))


def video_rows(videos, comments, info_by_lang, weeks, by="video"):
    by_v = defaultdict(list)
    for c in comments:
        by_v[c["vid"]].append(c)
    rows, cross = [], []
    for rank, v in enumerate(sorted(videos, key=lambda v: -(v["vpd"] or 0)), 1):
        cs = by_v.get(v["id"], [])
        likes = sum(c["likes"] for c in cs)
        cl = Counter(c["cluster"] for c in cs)
        cl_likes = Counter()
        for c in cs:
            cl_likes[c["cluster"]] += c["likes"]
        top = [f"{cluster_id(v['lang_group'], k)} {info_by_lang[v['lang_group']].get(k, {}).get('label', '')} — {n}개 ({n / len(cs):.0%})"
               for k, n in cl.most_common() if k != 0][:3]               # 내용 있는 묶음만 (0번 = 이모지·ㅋㅋ 는 따로 칸)
        best = max(cs, key=lambda c: c["likes"]) if cs else None
        wk = week_of_age(v["age_days"])
        rows.append({"순위(하루 조회수)": rank, "제목": v["title"], "링크": f"https://www.youtube.com/watch?v={v['id']}", "채널": v["channel"],
                     "구독자": v["subs"], "조회수": v["views"], "하루 조회수": v["vpd"], "올린 날": v["published"].date().isoformat() if v["published"] else "",
                     "주": f"{wk}주 전" if wk and wk <= weeks else "그 전", "언어": v["lang_group"], "나라": v["country"], "묶음(분석)": v["group"],
                     "시청자 댓글 수": len(cs), "좋아요 합": likes, "평균 좋아요": round(likes / len(cs), 2) if cs else None, "묶음 수": len(cl),
                     "낱말 없음(이모지·ㅋㅋ) 댓글 수": cl.get(0, 0),
                     "많은 묶음 1": top[0] if len(top) > 0 else "", "많은 묶음 2": top[1] if len(top) > 1 else "", "많은 묶음 3": top[2] if len(top) > 2 else "",
                     "좋아요 1위 댓글": f"👍{best['likes']:,} {clean(best['text'])}" if best else ""})
        if by == "comment":                                       # 댓글 단 날 기준 — 영상마다 주별 댓글 수·좋아요
            wk_n, wk_l = Counter(c["week"] for c in cs), Counter()
            for c in cs:
                wk_l[c["week"]] += c["likes"]
            for w in range(1, weeks + 1):
                rows[-1][f"{w}주 전 댓글 수"] = wk_n.get(w, 0)
                rows[-1][f"{w}주 전 좋아요 합"] = wk_l.get(w, 0)
        for k, n in cl.most_common():
            cross.append({"순위(하루 조회수)": rank, "제목": v["title"], "하루 조회수": v["vpd"], "주": rows[-1]["주"], "언어": v["lang_group"],
                          "묶음": cluster_id(v["lang_group"], k), "주제 낱말": info_by_lang[v["lang_group"]].get(k, {}).get("label", ""),
                          "댓글 수": n, "그 영상 댓글 중 몫": n / len(cs), "좋아요 합": cl_likes[k], "그 영상 좋아요 중 몫": cl_likes[k] / likes if likes else None})
    return rows, cross


# ---------------------------------------------------------------- 엑셀
def write_xlsx(path, sheets, guide_lines):
    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame({"안내": guide_lines}).to_excel(xw, sheet_name="안내", index=False)
        for name, rows in sheets:
            pd.DataFrame(rows).to_excel(xw, sheet_name=name[:31], index=False)
        wb = xw.book
        head_fill = PatternFill("solid", fgColor="DDEBF7")
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            if ws.max_row > 1 and ws.title != "안내":
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = head_fill
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for col in range(1, ws.max_column + 1):
                name = str(ws.cell(row=1, column=col).value or "")
                letter = get_column_letter(col)
                wide = any(x in name for x in ("댓글 1", "댓글 2", "댓글 3", "위 댓글", "제목", "안내", "많은 묶음", "주제 낱말"))
                ws.column_dimensions[letter].width = 60 if wide and ("댓글" in name or name == "안내") else (40 if wide else max(10, min(18, len(name) * 2 + 2)))
                if "몫" in name:
                    for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
                        for c in cell:
                            c.number_format = "0.0%"
                elif any(x in name for x in ("수", "합", "조회수", "구독자")):
                    for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
                        for c in cell:
                            c.number_format = "#,##0"
                if name == "링크":
                    for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
                        for c in cell:
                            if c.value:
                                c.hyperlink = c.value
                                c.font = Font(color="0563C1", underline="single")
                if wide:
                    for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
                        for c in cell:
                            c.alignment = Alignment(wrap_text=True, vertical="top")
    log("XLSX", file=str(path), sheets=len(sheets) + 1, kb=round(path.stat().st_size / 1024))


# ---------------------------------------------------------------- 실행
def run(topic, weeks=4, by="video", k=None, split=SPLIT_SHARE, out_path=None, now=None, period_spec=""):
    now = now or datetime.now(timezone.utc)
    k = {**DEFAULT_K, **(k or {})}
    out, videos, comments = load(topic, now)
    vids = {v["id"]: v for v in videos}
    info_by_lang = {}
    for lg in LANGS:
        rows = [c for c in comments if c["lang_group"] == lg]
        if not rows:
            info_by_lang[lg] = {}
            continue
        labels, sim, info = cluster_lang(lg, rows, k[lg], split)
        for c, l, s in zip(rows, labels, sim):
            c["cluster"], c["sim"] = int(l), float(s)
        info_by_lang[lg] = info
    for c in comments:
        if by == "comment":
            c["week"] = week_of_age((now - c["published"]).total_seconds() / 86400) if c["published"] else None
        else:
            c["week"] = week_of_age(vids[c["vid"]]["age_days"])
    ps = parse_periods(period_spec, weeks)
    weeks = max([weeks] + [b for _, _, b in ps if b])          # 영상별 '주' 칸 = 가장 먼 기간까지
    log("PERIOD", periods=" | ".join(p[0] for p in ps), by=by)
    sheets, compare = [], {}
    # 묶음 목록 (전체)
    all_rows = stats_rows(comments, info_by_lang, typical_n=REP_TYPICAL, liked_n=5)
    sheets.append(("묶음 목록(전체)", all_rows))
    per = {}
    for p in ps:
        sel = [c for c in comments if in_period(c["week"], p)]
        per[p[0]] = stats_rows(sel, info_by_lang)
        log("PERIOD", name=p[0], comments=len(sel), likes=sum(c["likes"] for c in sel), clusters=len(per[p[0]]))
    # 기간 비교 (댓글 수·좋아요 — 묶음마다 한 줄, 기간마다 열)
    keys = [(r["언어"], r["묶음"], r["주제 낱말"], r["대표 댓글 1"]) for r in all_rows]
    hits = {p[0]: {(r["언어"], r["묶음"]): r for r in per[p[0]]} for p in ps}
    for what, val, share in (("댓글 수", "댓글 수", "댓글 몫"), ("좋아요", "좋아요 합", "좋아요 몫")):
        rows = []
        for lg, cl, label, rep in keys:
            row = {"언어": lg, "묶음": cl, "주제 낱말": label}
            for p in ps:
                hit = hits[p[0]].get((lg, cl))
                row[f"{p[0]} {val}"] = hit[val] if hit else 0
                row[f"{p[0]} {share}"] = hit[share] if hit else 0.0
            row["대표 댓글 1"] = rep
            rows.append(row)
        compare[what] = rows
    sheets.append(("기간 비교(댓글 수)", compare["댓글 수"]))
    sheets.append(("기간 비교(좋아요)", compare["좋아요"]))
    for p in ps:
        sheets.append((f"기간 {p[0]}", per[p[0]]))
    vrows, cross = video_rows(videos, comments, info_by_lang, weeks, by)
    sheets.append(("영상별(하루 조회수 순)", vrows))
    sheets.append(("영상×묶음", cross))
    out_path = Path(out_path) if out_path else out / "comment_clusters.xlsx"
    guide = [
        f"자동 생성 — tools/comment_clusters.py {VERSION} ({now.isoformat(timespec='minutes')}). 연구·교육용 — 투자 권유가 아님.",
        f"자료: 승인한 영상 {len(videos):,}개의 시청자 댓글 {len(comments):,}개 (채널 주인·빈 댓글 뺌, 답글 포함). 언어 묶음 = 영상 언어 (한국어·영어·그 밖).",
        "묶기: 언어마다 댓글의 내용 낱말(한국어 = Kiwi 형태소 분석의 명사·동사·형용사 원형, 영어 = 낱말, 그 밖 = 낱말·두 글자 조각)로 TF-IDF → SVD "
        f"{SVD_DIM}차원 → MiniBatchKMeans (처음 묶음 수 {', '.join(f'{lg} {k[lg]}' for lg in LANGS)}; 한 묶음이 그 언어 댓글의 {split:.0%} 를 넘으면 다시 나눔; 씨앗 {SEED} — 다시 돌려도 같음).",
        "'비슷하다' = 같은 낱말을 쓴다는 뜻. 뜻이 같아도 다른 낱말로 쓰면 다른 묶음에, 같은 낱말을 다른 뜻으로 쓰면 같은 묶음에 들 수 있다 — 대표 댓글을 읽고 판단할 것.",
        f"묶음 번호 KO-001·EN-001·OT-001 = 그 언어에서 댓글이 많은 순. -000 = {EMPTY}.",
        "주제 낱말 = 그 묶음에서 다른 댓글보다 유난히 많이 쓴 낱말 5개 (묶음 TF-IDF 평균 − 전체 평균). 대표 댓글 = 묶음 한가운데에 가장 가까운 댓글(8자 이상). "
        "좋아요 n위 댓글 = 그 기간·묶음에서 한가운데에 가까운 절반 안의 좋아요 많은 댓글 (낱말 하나만 겹쳐 들어온 댓글이 보기로 나오지 않게 — 좋아요 합·몫은 묶음 전부).",
        f"기간 기준: {'댓글 단 날' if by == 'comment' else '영상 올린 날'} — 1주 전 = 0~6일 전, 2주 전 = 7~13일 전 … · 1~N주 전 = 누적 · 전체 = 기간 상관없이 모두. "
        f"이번 기간: {' | '.join(p[0] for p in ps)}. 시트 이름 '기간 1주 전' 등. 몫 = 그 기간·그 언어 안에서의 비율.",
        "기간 바꾸기: --weeks N (1~N주 전 한 주씩 + 1~2…1~N 누적) · --periods \"1,2,1-3,2-4\" (직접 — 2-4 = 2~4주 전) · --by comment (댓글 단 날 기준).",
        "기간 비교 시트: 묶음마다 한 줄, 기간마다 수(또는 좋아요 합)와 몫 — 같은 묶음이 기간에 따라 늘고 주는지 본다.",
        "영상별 시트: 하루 조회수(영상 정보를 받은 때 조회수 ÷ 올린 날부터 그때까지 날수, 최소 1일) 큰 순. 묶음(분석) = analysis.md 의 성공 사례 · 하루 조회수 상위(1-3번) · 그 밖. "
        "많은 묶음 1~3 = 내용 있는 묶음 중 댓글 많은 순 (-000 은 '낱말 없음' 칸에 따로).",
        "영상×묶음 시트: 영상마다 묶음별 댓글 수·좋아요 — 피벗 표로 다시 묶을 수 있다."
        + (" 댓글 단 날 기준이라 영상별 시트에 주마다 댓글 수·좋아요 칸이 있다." if by == "comment" else ""),
        "한계: 댓글·좋아요는 받은 때(승인한 날)까지의 값 — 올린 지 얼마 안 된 영상(1주 전)은 댓글이 덜 쌓였다. 큰 채널 댓글이 좋아요가 많다(좋아요 몫은 큰 채널 쪽으로 기운다). 외국(영어 등) 댓글 통계로 한국어 규칙을 만들지 않는다 (CLAUDE.md).",
        f"다시 만들기: python3 tools/comment_clusters.py {topic} --by {by} " + (f"--periods \"{period_spec}\"" if period_spec else f"--weeks {weeks}"),
    ]
    write_xlsx(out_path, sheets, guide)
    write_md(out_path.with_suffix(".md"), per, ps, info_by_lang, vrows, len(comments), by)
    return {"xlsx": str(out_path), "md": str(out_path.with_suffix(".md")), "comments": len(comments), "clusters": {lg: len(info_by_lang[lg]) for lg in LANGS}}


def write_md(md_path, per, ps, info_by_lang, vrows, n_comments, by):
    """깃허브에서 바로 보는 짧은 요약 — 자세한 것은 엑셀."""
    L = ["# 비슷한 댓글 묶음 — 요약 (자세한 것은 `comment_clusters.xlsx`)", "",
         f"`tools/comment_clusters.py` {VERSION} · 시청자 댓글 {n_comments:,}개 · 기간 기준 {'댓글 단 날' if by == 'comment' else '영상 올린 날'}. "
         "※ 연구·교육용 — 투자 권유가 아님. '비슷하다' = 같은 낱말을 쓴다는 뜻 (엑셀 '안내' 시트).", ""]
    for p in ps:
        rows = [r for r in per[p[0]] if r["언어"] == "한국어"][:8]
        if not rows:
            continue
        L += [f"## 한국어 · {p[0]} — 댓글 많은 묶음 8개", "", "| 묶음 | 주제 낱말 | 댓글 수 (몫) | 좋아요 합 (몫) | 대표 댓글 | 좋아요 1위 댓글 |", "|---|---|---|---|---|---|"]
        for r in rows:
            L.append(f"| {r['묶음']} | {r['주제 낱말']} | {r['댓글 수']:,} ({r['댓글 몫']:.0%}) | {r['좋아요 합']:,} ({r['좋아요 몫']:.0%}) | "
                     f"{clean(r.get('대표 댓글 1', ''), 60).replace('|', '∣')} | {clean(r.get('좋아요 1위 댓글', ''), 80).replace('|', '∣')} |")
        L.append("")
    L += ["## 하루 조회수 상위 영상 10개 — 많은 묶음", "", "| 순위 | 제목 | 하루 조회수 | 댓글 | 많은 묶음 1 |", "|---|---|---|---|---|"]
    for r in vrows[:10]:
        L.append(f"| {r['순위(하루 조회수)']} | [{clean(r['제목'], 50).replace('|', '∣')}]({r['링크']}) | {r['하루 조회수'] or 0:,.0f} | {r['시청자 댓글 수']:,} | "
                 f"{clean(r['많은 묶음 1'], 70).replace('|', '∣')} |")
    md_path.write_text("\n".join(L) + "\n", "utf-8")                      # 엑셀 옆에 (--out 이면 그 폴더)


def parse_k(s):
    res = {}
    for part in (s or "").split(","):
        if "=" in part:
            a, b = part.split("=", 1)
            a = {"그밖": "그 밖 언어", "그 밖": "그 밖 언어", "ko": "한국어", "en": "영어", "other": "그 밖 언어"}.get(a.strip(), a.strip())
            res[a] = int(b)
    return res


def main(argv=None):
    ap = argparse.ArgumentParser(description="내용이 비슷한 댓글 묶음 — 기간별·영상별 통계 엑셀 (할당량 0)")
    ap.add_argument("topic")
    ap.add_argument("--weeks", type=int, default=4, help="한 주씩·누적 기간을 몇 주 전까지 (기본 4)")
    ap.add_argument("--periods", default="", help='기간 직접 지정, 예: "1,2,1-3,2-4" (2-4 = 2~4주 전) — 주면 --weeks 대신')
    ap.add_argument("--by", choices=["video", "comment"], default="video", help="기간 기준: video = 영상 올린 날(기본) · comment = 댓글 단 날")
    ap.add_argument("--k", default="", help="처음 묶음 수, 예: 한국어=60,영어=40,그밖=15")
    ap.add_argument("--split", type=float, default=SPLIT_SHARE, help="한 묶음이 그 언어 댓글의 이 몫을 넘으면 다시 나눔 (기본 0.03)")
    ap.add_argument("--out", default="", help="엑셀 파일 경로 (기본 <주제>/result/outliers/comment_clusters.xlsx)")
    ap.add_argument("--log-level", default="INFO", choices=["DEBUG", "INFO", "WARNING"])
    a = ap.parse_args(argv)
    logging.basicConfig(level=getattr(logging, a.log_level), format="%(message)s", stream=sys.stderr)
    log("START", topic=a.topic, weeks=a.weeks, periods=a.periods or "기본", by=a.by, k=a.k or "기본", split=a.split, tool=f"comment_clusters.py {VERSION}")
    t0 = time.time()
    res = run(a.topic, weeks=a.weeks, by=a.by, k=parse_k(a.k), split=a.split, out_path=a.out or None, period_spec=a.periods)
    log("DONE", seconds=round(time.time() - t0, 1), **{k: v for k, v in res.items() if k != "clusters"}, clusters=json.dumps(res["clusters"], ensure_ascii=False))


if __name__ == "__main__":
    main()
