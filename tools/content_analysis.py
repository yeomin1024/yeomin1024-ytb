#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# VERSION: v1.0.2 — 2026-10-28 — 엑셀 '읽는 법'의 썸네일 장 수·성공 사례 뜻을 주제 파일에서 (주식 167장이 굳어 있어 역사·재테크에 틀린 수가 찍혔음).
# (v1.0.1 — 2026-10-28 — 썸네일을 지금 묶음 기준으로 (다른 세션이 성공 기준을 고쳐 158 → 114개 — 통계는 지금 성공 사례만, 영상별 시트에 묶음·지금 번호·분류 때 번호).
# (v1.0 — 2026-10-28 — 최초.) 사용자 지시 "아래처럼 분석하고 있는지 확인하고 아니면 반영해 그리고 내가 확인할 수 있는 엑셀 생성해 —
#          제목: 문구 유형, 사용 단어 · 썸네일: 이미지, 주요 문구, 보조 문구 유형 · 대본 내용: 제목과 도입부 연결, 전개 방식, 말투, 정보를 알려주는 방식,
#          흥미 유발 · 댓글: 내용별 유형, 공감 많은 댓글, 어떤 부분이 좋았는지, 아쉬운 피드백".
#          승인한 영상 폴더(W_·01_·P01_)를 영상마다 네 갈래로 나눠 엑셀 한 권에 — 성공 사례·하루 조회수 상위·그 밖 승인 영상을 칸으로 견줌.
#
# CHANGELOG
# | 버전 | 날짜 | 바뀐 것 | 함수 |
# |---|---|---|---|
# | v1.0.2 | 2026-10-28 | 읽는 법 — 썸네일 분류 장 수 = thumb_codes.tsv 줄 수 · 성공 사례 뜻 = 주제 규칙 | GUIDE · write_xlsx · main |
# | v1.0.1 | 2026-10-28 | 썸네일 통계 = 지금 성공 사례만 · 묶음·지금 번호 칸 | thumb_rows · thumb_stats |
# | v1.0 | 2026-10-28 | 최초 | load · title_types · thumb_rows · script_rows · comment_rows · write_xlsx · write_md |
#
# 실행: python3 tools/content_analysis.py stock            → stock/result/outliers/content_analysis.xlsx · content_analysis.md
#       python3 tools/content_analysis.py stock --debug    (DEBUG 로그)
# 입력: <주제>/result/outliers/ 의 analysis.json(성공 사례·하루 조회수 상위 목록) · 영상 폴더(metadata·comments.csv·transcript.json·thumbnail.jpg)
#       · thumb_codes.tsv(성공 사례 썸네일 — Claude 가 그림을 보고 분류) — 네트워크 없음.
# 규칙은 모두 낱말 규칙(정규식)이다 — 한국어 영상에 맞춰 만들었고, 외국어 영상은 숫자·물음표 정도만 잡힌다. 댓글 유형은 tools/comment_types.py 규칙 v1 그대로
# (사람 표본 시험 정확도 약 45% — comment_types.md 2번) → 몫은 대강의 크기로만 읽는다.
# ※ 연구·교육용.
import csv
import json
import logging
import math
import re
import statistics as st
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import comment_types as ct  # noqa: E402

VERSION = "v1.0.2"
log = logging.getLogger("content_analysis")


class StageFormatter(logging.Formatter):
    def format(self, r):
        return f"[{getattr(r, 'stage', 'RUN')}] [{self.formatTime(r, '%Y-%m-%dT%H:%M:%S')}] {r.getMessage()}"


def out(stage, level=logging.INFO, **kv):
    log.log(level, " ".join(f"{k}={v}" for k, v in kv.items()), extra={"stage": stage})


F = re.I
# ---------------------------------------------------------------- 1. 제목·문구 유형 (한 제목에 여러 유형 — 위에서 아래 순서는 뜻 없음)
TITLE_TYPES = [
    ("질문형", r"\?|？|까요|나요|을까|ㄹ까|할까|일까|인가|는가|됐나|되나|하나요"),
    ("숫자형", r"\d"),
    ("손해·경고형", r"손해|손실|잃|날리|날렸|날립|폭락|붕괴|망하|망함|망했|위험|함정|경고|주의|절대|하지\s?마|사지\s?마|말라|조심|무너|털리|털렸|물린|물렸|피하|"
                r"loss|crash|trap|warning|don'?t|never|avoid|mistake|exposed|downside"),
    ("이유·방법형", r"이유|방법|하는\s?법|는법|비법|매매법|투자법|원칙|전략|순서|공식|how\s|why\s|ways|tips|guide|strategy"),
    ("비교형", r"\bvs\b|보다|비교|차이|승자|누가 이길|도 아닌|\bor\b|than\b|better"),
    ("경험·고백형", r"했습니다|했어요|봤습니다|봤더니|해보니|해봤|었더니|았더니|떠납니다|안합니다|샀다|팔고|제가|내가|나는|\bI\b|I'm|I just|my\s"),
    ("권위·인물형", r"교수|전문가|박사|대표|회장|애널리스트|펀드매니저|소장|회계사|자산가|부자|천재|트레이더|CEO|버핏|expert|billionaire|millionaire|rich people|trader"),
    ("결과·실험형", r"결과|실험|백테스트|해보니|했더니|년\s?뒤|년이면|만들었|만든|됐|되었|backtest|tested|results?\b|after \d+|years? of"),
    ("시점·긴급형", r"지금|오늘|당장|곧\s|내일|이번|긴급|속보|임박|2026|2025|now\b|today|this week"),
    ("목록형", r"\d+\s*(가지|개|종목|명|단계)|TOP\s?\d|top \d|\d+\s*(stocks|etfs|funds|ways|things)"),
    ("인용형", r"[\"“”‘’']"),
    ("호기심·비밀형", r"진짜|사실|비밀|아무도|숨겨진|몰랐|소름|충격|놀라|반전|수상|심상치|정체|secret|nobody|truth|hidden|surprising"),
    ("감탄·강조형", r"!|！|\.\.|…|\?\?"),
]
TITLE_RX = [(n, re.compile(p, F)) for n, p in TITLE_TYPES]
JOSA = re.compile(r"(으로|에서|에게|까지|부터|이랑|하고|처럼|보다|만큼|은|는|이|가|을|를|의|에|도|로|와|과|만|요)$")
STOP = {"그리고", "그냥", "이것", "이거", "그것", "정말", "너무", "the", "a", "an", "of", "to", "and", "in", "is", "for", "on", "it", "you", "your", "this", "with"}


def title_types(text):
    return [n for n, rx in TITLE_RX if rx.search(text or "")]


def words(text):
    """제목 낱말 — 한글은 끝 조사를 거칠게 떼고, 영문은 소문자. 한 글자·불용어 뺌."""
    ws = []
    for w in re.findall(r"[가-힣]+|[A-Za-z][A-Za-z0-9&+.]*|\d+[%만억원배년개]*", text or ""):
        if re.match(r"[가-힣]", w):
            w2 = JOSA.sub("", w)
            w = w2 if len(w2) >= 2 else w
        else:
            w = w.lower().rstrip(".")
        if len(w) >= 2 and w not in STOP:
            ws.append(w)
    return ws


# ---------------------------------------------------------------- 2. 썸네일 보조 문구 유형
SUB_TYPES = [
    ("말풍선·인용", r"[\"“”]|말풍선"),
    ("출연자 표시", r"이름표|대표|소장|교수|회장|부사장|단장|과장|평론가|기자|박사|이름|Ft\.|feat"),
    ("로고·상표", r"로고|SAMSUNG|SK hynix|KRX|Schwab|Vanguard|YTN|PODCAST|Wealth"),
    ("숫자·수익률", r"\d"),
    ("질문", r"\?|？"),
    ("꼬리표·강조", r"꼬리표|Now|Hot|속보|필수|추천|풀영상|통합본|✕"),
]
SUB_RX = [(n, re.compile(p, F)) for n, p in SUB_TYPES]

# ---------------------------------------------------------------- 3. 대본 (한국어 규칙)
FIRST_TYPES = [
    ("인사·소개", r"^.{0,12}(안녕하세요|여러분|반갑습니다|오늘은|오늘 영상)"),
    ("질문으로 시작", r"까요|나요|을까|할까|인가요|[?？]|궁금"),
    ("숫자로 시작", r"\d|억|만 원|퍼센트|%"),
    ("사연·경험으로 시작", r"했어요|했습니다|였어요|었어요|았어요|제가|저는|저희|결국|그때"),
    ("결론·주장으로 시작", r"입니다|합니다|됩니다|없습니다|있습니다|해야|하세요"),
]
FIRST_RX = [(n, re.compile(p)) for n, p in FIRST_TYPES]
RX_NUM = re.compile(r"\d+(?:[.,]\d+)?\s*(?:%|퍼센트|배|억|만|원|달러|년|개월|주)?")
RX_CMP = re.compile(r"보다|비교|차이|대비|반면|에 비해|vs")
RX_EX = re.compile(r"예를 들어|예를 들면|예컨대|가령|만약|만일|이를테면")
RX_EVID = re.compile(r"연구|자료|통계|데이터|보고서|발표|조사|기사|공시|백테스트|논문")
RX_Q = re.compile(r"까요|나요|을까|ㄹ까|할까|인가요|일까요|아시나요|[?？]")
RX_TEASE = re.compile(r"잠시 후|뒤에서|끝까지|마지막에|이따가|나중에|곧 말씀|그런데|하지만|그러나|반전|놀랍게도|진짜 이유|비밀|중요한 건|핵심은")
RX_LIST = re.compile(r"첫째|첫 번째|둘째|두 번째|셋째|세 번째|1번|2번|3번|마지막으로")
RX_STORY = re.compile(r"제가|저는|저희|했어요|했습니다|였어요|그때|당시")
END_FORMAL = re.compile(r"(니다|니까|십시오)$")
END_POLITE = re.compile(r"(요|죠)$")
END_PLAIN = re.compile(r"(다|어|야|지|냐|네|자|거든|잖아)$")

# ---------------------------------------------------------------- 4. 댓글 — 좋았던·아쉬운 부분 갈래
ASPECTS = [
    ("설명·이해", r"설명|이해|쉽게|알기 쉽|귀에 쏙|명쾌|explain|clear"),
    ("정리·요약", r"정리|요약|핵심|summary"),
    ("숫자·자료·계산", r"숫자|데이터|자료|통계|백테스트|계산|그래프|차트|수익률|data|numbers|chart"),
    ("현실·공감·솔직", r"현실적|현실을|공감|솔직|진정성|허황|honest|realistic"),
    ("사례·경험", r"사례|경험|이야기|사연|썰|story"),
    ("구성·편집·자막", r"편집|자막|구성|화면|영상미|edit"),
    ("목소리·말투", r"목소리|말투|발음|음성|ai ?목소리|voice"),
    ("길이·반복", r"너무 길|길어|길다|길게|반복|지루|늘어지|too long|boring|repetitive"),
    ("제목·썸네일·낚시", r"제목|썸네일|어그로|낚시|clickbait|title|thumbnail"),
    ("광고·홍보", r"광고|협찬|홍보|sponsor|\bad\b"),
    ("정확성·오류", r"틀렸|틀린|잘못|오류|거짓|사실과|정정|wrong|mistake|incorrect"),
    ("출연자·진행자", r"출연|진행자|선생님|교수님|대표님|님 ?말씀|host|guest"),
]
ASPECT_RX = [(n, re.compile(p, F)) for n, p in ASPECTS]
# 아쉬운 부분 = 영상 자체에 대한 피드백 낱말만 (일반 투자 의견·정치 비판이 '비판'으로 섞이는 것을 막음 — v1.0 점검에서 고침)
BAD_ASPECTS = [
    ("길이·반복", r"너무 길|길어|길다|반복|지루|늘어지|too long|boring|repetitive|dragged"),
    ("제목·썸네일·낚시", r"제목|썸네일|어그로|낚시|clickbait|title|thumbnail"),
    ("광고·홍보", r"광고|협찬|홍보|sponsor|\bads?\b"),
    ("정확성·오류", r"틀렸|틀린|잘못 (말|알|설명)|오류|거짓|사실과 다|정정|wrong|mistake|incorrect|misleading"),
    ("목소리·말투", r"목소리|말투|발음|음성|voice|accent"),
    ("구성·편집·소리", r"편집|자막|배경음|음악|소리가|화면이|edit|music|audio"),
    ("설명 부족·어려움", r"어렵|이해가 안|설명이 부족|헷갈|무슨 말|confusing|unclear"),
    ("결론·근거 부족", r"결론이 없|근거가 없|근거 없|출처|요점이|뭘 말하|no source|what'?s the point"),
]
BAD_RX = [(n, re.compile(p, F)) for n, p in BAD_ASPECTS]
RX_PRAISE_EN = re.compile(r"\bi love|great (video|job|content)|thank(s| you)|appreciate|well explained", F)
RX_BOT = re.compile(r"likes the comment|1k likes|whatsapp|telegram|wa\.me", F)
RX_TIME = re.compile(r"(?<!\d)(\d{1,2}):([0-5]\d)(?!\d)")
RX_REQ = re.compile(r"해\s?주세요|해주시면|다뤄|알려\s?주|부탁|원합니다|please|could you|can you", F)


def aspects(text):
    return [n for n, rx in ASPECT_RX if rx.search(text or "")]


def med(xs):
    xs = [x for x in xs if x is not None]
    return round(st.median(xs), 2) if xs else None


def share(k, n):
    return round(k / n * 100, 1) if n else None


def short(t, n=160):
    t = re.sub(r"\s+", " ", str(t or "")).strip()
    return t if len(t) <= n else t[:n - 1] + "…"


# ---------------------------------------------------------------- 읽기
def load(topic):
    od = ROOT / topic / "result" / "outliers"
    try:
        a = json.loads((od / "analysis.json").read_text("utf-8"))
    except OSError as e:
        out("LOAD", logging.ERROR, file="analysis.json", error=e, next_step=f"python3 tools/cloud_run.py {topic} analyze 먼저")
        raise
    succ = {c["id"]: c for c in a.get("success_cases", [])}
    vpd = {c["id"]: c for c in (a.get("vpd_top") or {}).get("cases", [])}
    codes = {}
    tp = od / "thumb_codes.tsv"
    if tp.is_file():
        with open(tp, encoding="utf-8") as f:
            for r in csv.DictReader((l for l in f if not l.startswith("#")), delimiter="\t"):
                codes[r["id"]] = r
    vids = []
    for d, meta, info, cms in ct.iter_videos(od):
        vid = meta.get("id")
        tj = d / "transcript.json"
        tr = json.loads(tj.read_text("utf-8")) if tj.is_file() else None
        grp = "성공 사례" if vid in succ else "하루 조회수 상위" if vid in vpd else "그 밖 승인 영상"
        vids.append(dict(dir=d, meta=meta, cms=cms, tr=tr, grp=grp, vpd=vid in vpd, no=succ.get(vid, {}).get("no")))
    out("LOAD", folder=od.relative_to(ROOT), videos=len(vids), success=len(succ), vpd_top=len(vpd), thumb_codes=len(codes),
        transcripts=sum(1 for v in vids if v["tr"]), comments=sum(len(v["cms"]) for v in vids))
    return od, a, vids, codes


# ---------------------------------------------------------------- 영상·제목
def video_rows(vids):
    rows = []
    for v in vids:
        m = v["meta"]
        rows.append({"영상 ID": m.get("id"), "묶음": v["grp"], "하루 조회수 상위에도": "예" if v["vpd"] and v["grp"] == "성공 사례" else "",
                     "성공 사례 번호": v["no"], "제목": m.get("title"), "채널": m.get("channel"), "언어": "외국어" if m.get("foreign") else (m.get("video_lang") or ""),
                     "구독자": m.get("subs"), "조회수": m.get("views"), "배수(조회수÷구독자)": m.get("ratio"), "하루 조회수": m.get("views_per_day"),
                     "길이(분)": m.get("minutes"), "올린 날": (m.get("published") or "")[:10], "대본": "있음" if v["tr"] else "",
                     "썸네일 분류": "있음" if v["no"] else "", "주소": m.get("url")})
    return rows


def title_rows(vids):
    rows = []
    for v in vids:
        t = v["meta"].get("title") or ""
        ty = title_types(t)
        rows.append({"영상 ID": v["meta"].get("id"), "묶음": v["grp"], "언어": "외국어" if v["meta"].get("foreign") else "한국어", "제목": t,
                     "글자 수": len(t), "문구 유형": " · ".join(ty) or "(없음)", "유형 수": len(ty), "낱말": " ".join(words(t)),
                     "배수": v["meta"].get("ratio"), "하루 조회수": v["meta"].get("views_per_day")})
    return rows


def title_stats(trows):
    ko = [r for r in trows if r["언어"] == "한국어"]
    g = defaultdict(list)
    for r in ko:
        g[r["묶음"]].append(r)
    S, V, O = g["성공 사례"], g["하루 조회수 상위"], g["그 밖 승인 영상"]
    rows = []
    for name, _ in TITLE_TYPES:
        has = lambda rs: [r for r in rs if name in r["문구 유형"].split(" · ")]  # noqa: E731
        rows.append({"문구 유형": name, "성공 사례 몫 %": share(len(has(S)), len(S)), "하루 조회수 상위 몫 %": share(len(has(V)), len(V)),
                     "그 밖 몫 %": share(len(has(O)), len(O)),
                     "성공 − 그 밖 (%p)": round((share(len(has(S)), len(S)) or 0) - (share(len(has(O)), len(O)) or 0), 1),
                     "이 유형 제목의 배수 중앙값": med([r["배수"] for r in has(ko)]), "이 유형 제목의 하루 조회수 중앙값": med([r["하루 조회수"] for r in has(ko)]),
                     "제목 수(한국어 전체)": len(has(ko))})
    rows.sort(key=lambda r: -r["성공 − 그 밖 (%p)"])
    base = {"성공 사례": len(S), "하루 조회수 상위": len(V), "그 밖 승인 영상": len(O)}
    lens = {k: med([r["글자 수"] for r in rs]) for k, rs in (("성공 사례", S), ("하루 조회수 상위", V), ("그 밖 승인 영상", O))}
    return rows, base, lens


def word_stats(trows, min_total=5):
    cnt = {k: Counter() for k in ("성공 사례", "하루 조회수 상위", "그 밖 승인 영상")}
    n = Counter()
    for r in trows:
        if r["언어"] != "한국어":
            continue
        n[r["묶음"]] += 1
        for w in set(r["낱말"].split()):
            cnt[r["묶음"]][w] += 1
    allw = Counter()
    for c in cnt.values():
        allw.update(c)
    rows = []
    for w, tot in allw.items():
        if tot < min_total:
            continue
        s, o = cnt["성공 사례"][w], cnt["그 밖 승인 영상"][w]
        rows.append({"낱말": w, "전체 제목 수": tot, "성공 사례 제목 수": s, "성공 사례 몫 %": share(s, n["성공 사례"]),
                     "하루 조회수 상위 몫 %": share(cnt["하루 조회수 상위"][w], n["하루 조회수 상위"]), "그 밖 몫 %": share(o, n["그 밖 승인 영상"]),
                     "성공 − 그 밖 (%p)": round((share(s, n["성공 사례"]) or 0) - (share(o, n["그 밖 승인 영상"]) or 0), 1)})
    rows.sort(key=lambda r: (-r["성공 − 그 밖 (%p)"], -r["전체 제목 수"]))
    return rows


# ---------------------------------------------------------------- 썸네일
def color_stats(path):
    """밝기·빨강·노랑 픽셀 몫 (64×36 으로 줄여 셈 — 글자색·배경색이 섞임, 대강의 값)."""
    try:
        from PIL import Image
        im = Image.open(path).convert("RGB").resize((64, 36))
    except Exception as e:  # noqa: BLE001 — 그림 하나가 깨져도 나머지는 계속
        out("THUMB", logging.WARNING, file=path, error=e, next_step="썸네일 다시 받기")
        return None, None, None
    raw = im.tobytes()
    px = [tuple(raw[i:i + 3]) for i in range(0, len(raw), 3)]
    n = len(px)
    bright = sum((r * 299 + g * 587 + b * 114) / 1000 for r, g, b in px) / n
    red = sum(1 for r, g, b in px if r > 150 and g < 90 and b < 90) / n * 100
    yel = sum(1 for r, g, b in px if r > 180 and g > 160 and b < 90) / n * 100
    return round(bright), round(red, 1), round(yel, 1)


def thumb_rows(vids, codes):
    rows = []
    for v in vids:
        vid = v["meta"].get("id")
        c = codes.get(vid)
        if not c:
            continue
        main, sub, title = c["main_text"], c["sub_text"], v["meta"].get("title") or ""
        mw, tw = set(words(main)), set(words(title))
        br, red, yel = color_stats(v["dir"] / "thumbnail.jpg")
        rows.append({"묶음": v["grp"], "성공 사례 번호(지금)": v["no"], "분류 때 번호": int(c["no"]), "영상 ID": vid, "제목": title, "이미지 유형": c["image"], "인물 표정": c["face"], "차트·그래프": "있음" if c["chart"] == "1" else "",
                     "주요 문구": main, "주요 문구 유형": " · ".join(title_types(main)) or "(없음)", "주요 문구 글자 수": len(main),
                     "보조 문구": sub, "보조 문구 유형": " · ".join(n for n, rx in SUB_RX if rx.search(sub)) or ("(없음)" if not sub else "그 밖"),
                     "제목과 같은 낱말 몫 %": share(len(mw & tw), len(mw)) if mw else None,
                     "밝기(0~255)": br, "빨강 몫 %": red, "노랑 몫 %": yel, "배수": v["meta"].get("ratio")})
    rows.sort(key=lambda r: (r["성공 사례 번호(지금)"] is None, r["성공 사례 번호(지금)"] or 0, r["분류 때 번호"]))
    return rows


def thumb_stats(trs):
    trs = [r for r in trs if r["묶음"] == "성공 사례"]  # 통계는 지금 성공 사례만 (분류는 됐지만 기준이 바뀌어 빠진 영상은 영상별 시트에만)
    n = len(trs)
    rows = []
    for col in ("이미지 유형", "인물 표정"):
        for k, c in Counter(r[col] for r in trs).most_common():
            rows.append({"항목": col, "값": k, "수": c, "몫 %": share(c, n), "배수 중앙값": med([r["배수"] for r in trs if r[col] == k])})
    c = sum(1 for r in trs if r["차트·그래프"])
    rows.append({"항목": "차트·그래프", "값": "있음", "수": c, "몫 %": share(c, n), "배수 중앙값": med([r["배수"] for r in trs if r["차트·그래프"]])})
    for col in ("주요 문구 유형", "보조 문구 유형"):
        cnt = Counter(t for r in trs for t in r[col].split(" · "))
        for k, c in cnt.most_common():
            rows.append({"항목": col, "값": k, "수": c, "몫 %": share(c, n), "배수 중앙값": med([r["배수"] for r in trs if k in r[col].split(" · ")])})
    ov = [r["제목과 같은 낱말 몫 %"] for r in trs if r["제목과 같은 낱말 몫 %"] is not None]
    rows.append({"항목": "주요 문구 ↔ 제목", "값": "같은 낱말 몫 중앙값 %", "수": len(ov), "몫 %": med(ov), "배수 중앙값": None})
    rows.append({"항목": "주요 문구 ↔ 제목", "값": "제목과 거의 같음(80% 이상)", "수": sum(1 for x in ov if x >= 80), "몫 %": share(sum(1 for x in ov if x >= 80), len(ov)), "배수 중앙값": None})
    rows.append({"항목": "주요 문구", "값": "글자 수 중앙값", "수": n, "몫 %": med([r["주요 문구 글자 수"] for r in trs]), "배수 중앙값": None})
    return rows


# ---------------------------------------------------------------- 대본
def text_between(snips, a, b):
    return " ".join(s["text"] for s in snips if a <= float(s["start"]) < b)


def endings(text):
    ws = re.findall(r"[가-힣]+", text)
    f = sum(1 for w in ws if END_FORMAL.search(w))
    p = sum(1 for w in ws if END_POLITE.search(w))
    pl = sum(1 for w in ws if END_PLAIN.search(w) and not END_FORMAL.search(w))
    return f, p, pl


def script_rows(vids):
    rows = []
    for v in vids:
        tr = v["tr"]
        if not tr:
            continue
        sn = [s for s in tr.get("snippets", []) if s.get("text")]
        if not sn:
            continue
        m = v["meta"]
        ko = str(tr.get("lang", "")).startswith("ko")
        dur = max(float(s["start"]) + float(s.get("duration", 0)) for s in sn)
        mins = max(dur / 60, 0.1)
        full = " ".join(s["text"] for s in sn)
        first15, first30, first60 = text_between(sn, 0, 8), text_between(sn, 0, 30), text_between(sn, 0, 60)  # 첫 문장 ≈ 첫 8초 (자동 자막엔 문장 부호가 없음)
        tw = set(words(m.get("title") or ""))
        hit = [w for w in tw if w in first60]
        quarters = []
        for i in range(4):
            seg = text_between(sn, dur * i / 4, dur * (i + 1) / 4)
            quarters.append(len(RX_Q.findall(seg)) + len(RX_TEASE.findall(seg)))
        f, p, pl = endings(full) if ko else (0, 0, 0)
        tot = f + p + pl
        tone = "외국어 — 원문 직접" if not ko else ("합니다체" if f >= max(p, pl) else "해요체" if p >= pl else "반말·해라체") if tot else ""
        ch = m.get("chapters") or []
        rows.append({"영상 ID": m.get("id"), "묶음": v["grp"], "제목": m.get("title"), "대본 언어": tr.get("lang"), "대본 길이(분)": round(dur / 60, 1),
                     "도입 30초": short(first30, 400),
                     "제목 낱말 중 첫 60초에 나온 몫 %": share(len(hit), len(tw)) if tw else None, "첫 60초에 나온 제목 낱말": " ".join(sorted(hit)),
                     "첫 문장 유형": next((n for n, rx in FIRST_RX if rx.search(first15)), "그 밖") if ko else "외국어",
                     "말투": tone, "합니다체 %": share(f, tot), "해요체 %": share(p, tot), "반말·해라체 %": share(pl, tot),
                     "숫자 /분": round(len(RX_NUM.findall(full)) / mins, 1), "비교 표현 /분": round(len(RX_CMP.findall(full)) / mins, 2) if ko else None,
                     "예시 표현 수": len(RX_EX.findall(full)) if ko else None, "근거 표현 수(연구·자료·통계…)": len(RX_EVID.findall(full)) if ko else None,
                     "사연·경험 표현 /분": round(len(RX_STORY.findall(full)) / mins, 2) if ko else None,
                     "질문 /분": round(len(RX_Q.findall(full)) / mins, 2), "예고·반전 표현 /분": round(len(RX_TEASE.findall(full)) / mins, 2) if ko else None,
                     "흥미 장치(질문+예고) 1/4": quarters[0], "2/4": quarters[1], "3/4": quarters[2], "4/4": quarters[3],
                     "흥미 장치 끝까지 이어짐": ("예" if min(quarters) > 0 else "아니오") if ko else "",
                     "목록형 전개 표현 수": len(RX_LIST.findall(full)) if ko else None,
                     "챕터 수": len(ch), "챕터": " / ".join(c.get("title", "") for c in ch)[:300],
                     "배수": m.get("ratio"), "하루 조회수": m.get("views_per_day")})
    return rows


def script_stats(srows):
    ko = [r for r in srows if str(r["대본 언어"]).startswith("ko")]
    cols = ["제목 낱말 중 첫 60초에 나온 몫 %", "합니다체 %", "해요체 %", "숫자 /분", "비교 표현 /분", "질문 /분", "예고·반전 표현 /분", "사연·경험 표현 /분", "챕터 수"]
    rows = []
    for g in ("성공 사례", "하루 조회수 상위", "그 밖 승인 영상"):
        rs = [r for r in ko if r["묶음"] == g]
        row = {"묶음": g, "한국어 대본 수": len(rs), "표본": "적음 — 차이로 규칙을 만들지 않음" if len(rs) < 20 else ""}
        for c in cols:
            row[c + " (중앙값)"] = med([r[c] for r in rs])
        row["흥미 장치 끝까지 이어진 몫 %"] = share(sum(1 for r in rs if r["흥미 장치 끝까지 이어짐"] == "예"), len(rs))
        for k, c in Counter(r["첫 문장 유형"] for r in rs).most_common(3):
            row.setdefault("첫 문장 유형 상위", "")
            row["첫 문장 유형 상위"] += f"{k} {c} · "
        rows.append(row)
    return rows


# ---------------------------------------------------------------- 댓글
def comment_rows(vids):
    per, liked_all, good_all, bad_all = [], [], [], []
    by_group = defaultdict(Counter)
    by_group_likes = defaultdict(Counter)
    asp_good, asp_bad = defaultdict(Counter), defaultdict(Counter)
    for v in vids:
        m = v["meta"]
        cs = [c for c in v["cms"] if not c["creator"] and c["text"].strip()]
        if not cs:
            per.append({"영상 ID": m.get("id"), "묶음": v["grp"], "제목": m.get("title"), "시청자 댓글 수": 0})
            continue
        tcount, good, bad, times, crit = Counter(), [], [], Counter(), 0
        for c in cs:
            ty = ct.classify(c["text"])
            c["type"] = ty
            tcount[ty] += 1
            by_group[v["grp"]][ty] += 1
            by_group_likes[v["grp"]][ty] += c["likes"]
            if RX_BOT.search(c["text"]):
                continue
            asp = aspects(c["text"])
            if ty in ("P", "I") and asp:
                good.append((c, asp))
                for a in asp:
                    asp_good[v["grp"]][a] += 1
            if ty == "C":
                crit += 1
            basp = [n for n, rx in BAD_RX if rx.search(c["text"])]
            if (ty == "C" or (ty == "Q" and RX_REQ.search(c["text"]))) and basp and not ct.PRAISE_STRONG.search(c["text"]) and not RX_PRAISE_EN.search(c["text"]):  # 영상의 부분을 짚은 비판·요청만
                bad.append((c, basp))
                for a in basp:
                    asp_bad[v["grp"]][a] += 1
            for mm, ss in RX_TIME.findall(c["text"]):
                times[int(mm)] += 1
        n = len(cs)
        top = sorted(cs, key=lambda c: -c["likes"])[:3]
        good.sort(key=lambda x: -x[0]["likes"])
        bad.sort(key=lambda x: -x[0]["likes"])
        row = {"영상 ID": m.get("id"), "묶음": v["grp"], "제목": m.get("title"), "시청자 댓글 수": n}
        for code in ct.TYPE_ORDER:
            row[ct.type_name(code) + " %"] = share(tcount[code], n)
        for i, c in enumerate(top, 1):
            row[f"공감 {i}위 (👍)"] = c["likes"]
            row[f"공감 {i}위 댓글"] = f"[{ct.type_name(c['type'])}] " + short(c["text"])
        row["좋았던 부분 댓글 수"] = len(good)
        row["좋았던 부분 갈래(상위)"] = " · ".join(f"{a} {k}" for a, k in Counter(a for _, asp in good for a in asp).most_common(3))
        row["좋았던 부분 대표 댓글"] = " ‖ ".join(f"👍{c['likes']} " + short(c["text"], 120) for c, _ in good[:2])
        row["비판·반박 댓글 수(전체)"] = crit
        row["아쉬운 부분·요청 댓글 수(영상의 부분을 짚은 것)"] = len(bad)
        row["아쉬운 부분 갈래(상위)"] = " · ".join(f"{a} {k}" for a, k in Counter(a for _, asp in bad for a in asp).most_common(3))
        row["아쉬운 부분 대표 댓글"] = " ‖ ".join(f"👍{c['likes']} " + short(c["text"], 120) for c, _ in bad[:2])
        row["시간 표시 댓글 수"] = sum(times.values())
        row["가장 많이 짚은 분(분:개)"] = " · ".join(f"{k}분:{c}" for k, c in times.most_common(3))
        per.append(row)
        for c in top:
            liked_all.append({"묶음": v["grp"], "영상 제목": m.get("title"), "👍": c["likes"], "유형": ct.type_name(c["type"]), "댓글": short(c["text"], 300), "영상 ID": m.get("id")})
        for c, asp in good[:5]:
            good_all.append({"묶음": v["grp"], "영상 제목": m.get("title"), "👍": c["likes"], "갈래": " · ".join(asp), "댓글": short(c["text"], 300), "영상 ID": m.get("id")})
        for c, asp in bad[:5]:
            bad_all.append({"묶음": v["grp"], "영상 제목": m.get("title"), "👍": c["likes"], "갈래": " · ".join(asp), "유형": ct.type_name(c["type"]),
                            "댓글": short(c["text"], 300), "영상 ID": m.get("id")})
    tstats = []
    for g in ("성공 사례", "하루 조회수 상위", "그 밖 승인 영상"):
        n, nl = sum(by_group[g].values()), sum(by_group_likes[g].values())
        for code in ct.TYPE_ORDER:
            tstats.append({"묶음": g, "유형": ct.type_name(code), "댓글 수": by_group[g][code], "몫 %": share(by_group[g][code], n),
                           "좋아요 합": by_group_likes[g][code], "좋아요 몫 %": share(by_group_likes[g][code], nl)})
    astats = []
    for kind, src in (("좋았던 부분", asp_good), ("아쉬운 부분·요청", asp_bad)):
        for g in ("성공 사례", "하루 조회수 상위", "그 밖 승인 영상"):
            tot = sum(src[g].values())
            for a, k in src[g].most_common():
                astats.append({"구분": kind, "묶음": g, "갈래": a, "댓글 수": k, "몫 %": share(k, tot)})
    keyf = lambda r: (-r["👍"])  # noqa: E731
    return per, tstats, astats, sorted(liked_all, key=keyf)[:300], sorted(good_all, key=keyf)[:300], sorted(bad_all, key=keyf)[:300]


# ---------------------------------------------------------------- 쓰기
GUIDE = [
    ("제목", "문구 유형(13가지 낱말 규칙 — 한 제목에 여러 개) · 사용 낱말(조사를 거칠게 뗌)", "제목_영상별 · 제목_유형통계 · 제목_낱말",
     "한국어 제목만 통계 · 묶음 비교는 '승인한 영상 안'의 비교 (공식 비교 집단은 analysis.md 4번·1-2번)"),
    ("썸네일", "이미지 유형·인물 표정·차트 · 주요 문구(가장 큰 글자)·보조 문구(작은 글자·꼬리표·말풍선)와 그 유형 · 제목과 같은 낱말 몫 · 밝기·빨강·노랑 몫",
     "썸네일_영상별 · 썸네일_통계", "Claude 가 그림을 보고 분류한 영상만 (thumb_codes.tsv — {thumb_n}장) · 통계는 지금 성공 사례만 · 색 몫은 줄인 그림으로 센 대강의 값"),
    ("대본", "제목 낱말이 첫 60초에 나오는 몫(연결) · 첫 문장 유형 · 말투(끝맺음) · 정보 방식(숫자·비교·예시·근거) · 흥미 장치(질문·예고)가 4구간에 고루 있는지 · 챕터",
     "대본_영상별 · 대본_통계", "대본이 있는 영상만 (자막·음성인식 — 대부분 없음, 보류 중) · 자동 자막은 문장 부호가 없어 끝맺음·질문 수는 대강"),
    ("댓글", "모든 시청자 댓글에 유형 하나(comment_types 규칙 v1) · 공감(좋아요) 상위 · 좋았던 부분(칭찬·정보 댓글 중 영상의 무엇을 짚었나 — 12갈래) · 아쉬운 부분·요청(비판·요청 중 영상의 부분을 짚은 것만) · 시간 표시(몇 분을 짚었나)",
     "댓글_영상별 · 댓글_유형통계 · 댓글_갈래통계 · 댓글_공감상위 · 댓글_좋았던점 · 댓글_아쉬운점", "유형 규칙 정확도 약 45%(사람 표본 시험) · 영어 댓글은 규칙이 덜 맞음"),
]


def write_xlsx(path, sheets, meta):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as e:
        out("WRITE", logging.ERROR, error=e, next_step="pip install openpyxl")
        raise
    wb = Workbook()
    ws = wb.active
    ws.title = "읽는 법"
    ws.append(["콘텐츠 분석 — " + meta["topic"], f"tools/content_analysis.py {VERSION} · {meta['generated']}"])
    ws.append(["영상", f"{meta['videos']}개 (성공 사례 {meta['S']} · 하루 조회수 상위 {meta['V']} · 그 밖 승인 영상 {meta['O']}) · 대본 {meta['tr']}개 · 시청자 댓글 {meta['cm']:,}개"])
    ws.append([])
    ws.append(["분석", "무엇을 보나", "시트", "한계"])
    for r in GUIDE:
        ws.append([x.format(**meta) for x in r])
    ws.append([])
    ws.append(["묶음 뜻", "성공 사례 = analysis.md 1번 (기준은 그 주제 analysis_rules.toml [success]) · 하루 조회수 상위 = 1-3번 (구독자·날짜 상관없이 상위 10%) · 그 밖 = 승인했지만 둘 다 아닌 영상"])
    ws.append(["※", "숫자·원문만 — 규칙을 만들 때는 analysis.md 0번 '한계'를 먼저 읽는다 (표본 적은 차이·대본 없는 사례로 규칙을 만들지 않음). 연구·교육용."])
    bold = Font(bold=True)
    for c in ws[4]:
        c.font = bold
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width, ws.column_dimensions["D"].width = 14, 80, 40, 60
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    head = PatternFill("solid", fgColor="DDEBF7")
    for name, rows in sheets:
        ws = wb.create_sheet(name)
        if not rows:
            ws.append(["(자료 없음)"])
            continue
        cols = list(dict.fromkeys(k for r in rows for k in r))
        ws.append(cols)
        for r in rows:
            ws.append([r.get(c) for c in cols])
        for c in ws[1]:
            c.font, c.fill = bold, head
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, c in enumerate(cols, 1):
            w = max(len(str(c)), *(len(str(r.get(c) or "")) for r in rows[:200]))
            ws.column_dimensions[get_column_letter(i)].width = min(max(8, w * 1.1), 70)
    wb.save(path)


def write_md(path, meta, tstat, base, lens, thstat, sstat, cstat):
    def top(rows, k=5):
        return " · ".join(f"{r['문구 유형']} {r['성공 사례 몫 %']}% (그 밖 {r['그 밖 몫 %']}%)" for r in rows[:k])
    img = [r for r in thstat if r["항목"] == "이미지 유형"][:3]
    face = [r for r in thstat if r["항목"] == "인물 표정"][:3]
    ov = next((r for r in thstat if r["값"] == "같은 낱말 몫 중앙값 %"), {})
    s = next((r for r in sstat if r["묶음"] == "성공 사례"), {})
    lines = [
        f"# 콘텐츠 분석 (제목·썸네일·대본·댓글) — {meta['topic']}", "",
        f"자동 생성 — `tools/content_analysis.py` {VERSION} ({meta['generated']}). 엑셀 `content_analysis.xlsx` 가 본문 — 여기는 한눈에 보기만. ※ 연구·교육용.", "",
        "## 0. 자료와 한계",
        f"- 영상 {meta['videos']}개 (성공 사례 {meta['S']} · 하루 조회수 상위 {meta['V']} · 그 밖 {meta['O']}) · 대본 {meta['tr']}개 · 시청자 댓글 {meta['cm']:,}개.",
        "- 모두 낱말 규칙 — 한국어에 맞춤. 댓글 유형 규칙 정확도 약 45%. 대본은 대부분 없음(자막 막힘·음성인식 보류) → 대본 칸은 표본 적음.", "",
        "## 1. 제목 (한국어)",
        f"- 제목 수: 성공 사례 {base['성공 사례']} · 하루 조회수 상위 {base['하루 조회수 상위']} · 그 밖 {base['그 밖 승인 영상']} · 글자 수 중앙값 {lens['성공 사례']} · {lens['하루 조회수 상위']} · {lens['그 밖 승인 영상']}",
        f"- 성공 사례에 더 많은 문구 유형: {top(tstat)}", "",
        f"## 2. 썸네일 (지금 성공 사례 중 분류된 {sum(1 for r in thstat if r['항목'] == '이미지 유형') and next((r['수'] for r in thstat if r['값'] == '글자 수 중앙값'), '')}장)",
        "- 이미지: " + " · ".join(f"{r['값']} {r['몫 %']}%" for r in img) + " · 표정: " + " · ".join(f"{r['값']} {r['몫 %']}%" for r in face),
        f"- 주요 문구와 제목의 같은 낱말 몫 중앙값 {ov.get('몫 %')}%", "",
        "## 3. 대본 (한국어 대본이 있는 영상만)",
        f"- 성공 사례 {s.get('한국어 대본 수')}개 — 제목 낱말 첫 60초 몫 중앙값 {s.get('제목 낱말 중 첫 60초에 나온 몫 % (중앙값)')}% · 숫자 {s.get('숫자 /분 (중앙값)')}/분 · 질문 {s.get('질문 /분 (중앙값)')}/분 · 첫 문장 {s.get('첫 문장 유형 상위', '')}", "",
        "## 4. 댓글 (성공 사례)",
        "- " + " · ".join(f"{r['유형']} {r['몫 %']}%" for r in sorted([r for r in cstat if r['묶음'] == '성공 사례'], key=lambda r: -(r['몫 %'] or 0))[:6]),
        "- 영상별 좋았던·아쉬운 부분·공감 상위·시간 표시는 엑셀 '댓글_영상별'·'댓글_좋았던점'·'댓글_아쉬운점'.", ""]
    path.write_text("\n".join(lines), "utf-8")


def main(argv=None):
    argv = argv or sys.argv[1:]
    if not argv:
        print(__doc__ or "사용: python3 tools/content_analysis.py <주제>")
        return 2
    topic = argv[0]
    h = logging.StreamHandler(sys.stdout)
    h.setFormatter(StageFormatter())
    log.addHandler(h)
    log.setLevel(logging.DEBUG if "--debug" in argv else logging.INFO)
    t0 = time.time()
    od, a, vids, codes = load(topic)
    vrows, trows = video_rows(vids), title_rows(vids)
    tstat, base, lens = title_stats(trows)
    wstat = word_stats(trows)
    out("TITLE", titles=len(trows), types=len(TITLE_TYPES), words_kept=len(wstat), top_type=tstat[0]["문구 유형"] if tstat else "-")
    thr = thumb_rows(vids, codes)
    thstat = thumb_stats(thr) if thr else []
    out("THUMB", coded=len(thr), coded_success_now=sum(1 for r in thr if r["묶음"] == "성공 사례"),
        missing_codes=sum(1 for v in vids if v["no"] and v["meta"].get("id") not in codes))
    srows = script_rows(vids)
    sstat = script_stats(srows)
    out("SCRIPT", transcripts=len(srows), korean=sum(1 for r in srows if str(r["대본 언어"]).startswith("ko")))
    per, cstat, astat, liked, good, bad = comment_rows(vids)
    out("COMMENT", videos=len(per), comments=sum(r.get("시청자 댓글 수", 0) for r in per), good=len(good), bad=len(bad))
    meta = dict(topic=topic, generated=time.strftime("%Y-%m-%d %H:%M"), videos=len(vids), S=sum(v["grp"] == "성공 사례" for v in vids),
                V=sum(v["grp"] == "하루 조회수 상위" for v in vids), O=sum(v["grp"] == "그 밖 승인 영상" for v in vids),
                tr=len(srows), cm=sum(r.get("시청자 댓글 수", 0) for r in per), thumb_n=len(codes))
    sheets = [("영상", vrows), ("제목_영상별", trows), ("제목_유형통계", tstat), ("제목_낱말", wstat), ("썸네일_영상별", thr), ("썸네일_통계", thstat),
              ("대본_영상별", srows), ("대본_통계", sstat), ("댓글_영상별", per), ("댓글_유형통계", cstat), ("댓글_갈래통계", astat),
              ("댓글_공감상위", liked), ("댓글_좋았던점", good), ("댓글_아쉬운점", bad)]
    xp, mp = od / "content_analysis.xlsx", od / "content_analysis.md"
    write_xlsx(xp, sheets, meta)
    write_md(mp, meta, tstat, base, lens, thstat, sstat, cstat)
    out("DONE", xlsx=xp.relative_to(ROOT), md=mp.relative_to(ROOT), sheets=len(sheets) + 1, sec=f"{time.time() - t0:.1f}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
