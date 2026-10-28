#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# VERSION: v1.1 — 2026-10-28 — 분석 고침(success_analysis v1.6·shorts_analysis v1.3 — 사용자 지시 "1은 그렇게하고 최근에 조회수가 높을수록 더 많이
#          반영되도록 해야해 3은 할당량 남으면 진행하도록 해 나머지는 너가 추천하는 방향으로 하고") 뒤: 분석 결과의 묶음(①# · ②V · ③T)과 경계를 그대로
#          읽어 판정 · ③ 모집단 = 기간 없는 검색 포함 · '어디서 봄' 에 기간 없는 검색 · 못 한 기간 없는 검색 수. 받은 때 하루 조회수 비교 칸은 뺌
#          (분석이 이제 받은 때 기준).
# v1.0 — 2026-10-28 — 새로: 분석 대상 세 묶음 확인 엑셀. 사용자 지시 2026-10-28 "아래 처럼 분석하고 있는지 확인하고 내가 확인할 수 있는
#          엑셀 파일 생성해 / 분석 대상(롱폼, 쇼츠 종합) — 구독자에 비해 업로드 대비 조회수가 높은 영상 (성공 사례) · 구독자 상관 없이 업로드 대비
#          조회수가 높은 영상 (최근 트렌드) · 업로드 기간 상관없이 조회수가 높은 영상 (최근 트렌드에 비슷한 과거 사례)".
#          · 분석 코드(success_analysis·shorts_analysis 의 analyze — 파일을 쓰지 않음)를 그대로 불러 같은 모집단·같은 규칙으로 셈 → 묶음마다
#            '기준 통과 / 자료 받음(승인) / 분석 묶음에 듦' 과 빠진 점. 지금 분석에 없는 묶음(쇼츠 하루 조회수 상위·조회수 상위)도 같은 꼴(상위 N%)로
#            세어 '분석 묶음 없음' 으로 보임 — 분석 결과(analysis.md)는 바꾸지 않음.
#          · 엑셀: 안내 · 확인 결과 · 모집단 · ①성공 사례 · ②최근 트렌드 · ③조회수 상위 · 겹침 · 전체 목록 (롱폼·쇼츠 한 표, '형식' 칸). 할당량 0.
#          ※ 연구·교육용 — 투자 권유가 아님.
"""분석 대상 세 묶음 확인 — 롱폼·쇼츠 종합 엑셀

  python3 tools/target_check.py stock            # stock/result/outliers/target_check.xlsx (+ .md 요약)
  python3 tools/target_check.py stock --pct 10   # ②·③ 상위 몇 % (기본 = analysis_rules.toml [vpd_top] TOP_PCT)
"""
import argparse
import json
import logging
import sys
import time
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import outlier_research as orr        # noqa: E402  top_pct_threshold (웹 화면 추천·분석 1-3번과 같은 경계)
import shorts_analysis as sha         # noqa: E402
import success_analysis as sa         # noqa: E402

VERSION = "v1.1"
ROOT = Path(__file__).resolve().parent.parent
LOG = logging.getLogger("tchk")
G1, G2, G3 = "① 성공 사례", "② 최근 트렌드", "③ 조회수 상위"
LONG, SHORT = "롱폼", "쇼츠"


def log(stage, level=logging.INFO, **kv):
    LOG.log(level, "[%s] [%s] %s", stage, datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S"),
            " ".join(f"{k}={json.dumps(v, ensure_ascii=False) if isinstance(v, str) and (' ' in v or not v) else v}" for k, v in kv.items()))


# ---------------------------------------------------------------- 자료 (분석 코드 그대로 — 파일은 쓰지 않음)
def load_long(topic, now):
    out = ROOT / topic / "result" / "outliers"
    rules, rel, themes = sa.load_rules(ROOT, topic), sa.load_rel(ROOT, topic), sa.load_themes(ROOT, topic)
    cache_infos, cache_subs = sa.load_cache(out)
    cases = sa.load_cases(out, cache_infos, cache_subs, now)
    mc = sa.load_market_caches(out)
    searched = {i for d in (mc[0] or ()) for i in (d.get("ids") or [])}
    alltime = sa.load_alltime_ids(out / ".cache") - searched                    # v1.1: 기간 없는 검색으로만 찾은 영상
    data, success, _group, rows_all, cases = sa.analyze(topic, topic, cases, cache_infos, cache_subs, rules, rel, themes, now, mc,
                                                        sa.load_uploads(out), alltime_ids=alltime)
    return {"rules": rules, "data": data, "success": success, "rows": rows_all, "cases": cases, "searched": searched, "alltime": alltime,
            "fetched": {k: v.get("fetched") for k, v in cache_infos.items()}}


def load_shorts(topic, now):
    out_all = ROOT / topic / "result" / "outliers"
    sdir, cache = sha.shorts_out(ROOT, topic), sha.shorts_cache(ROOT, topic)
    rules, rel, themes = sha.load_rules(ROOT, topic), sa.load_rel(ROOT, topic), sa.load_themes(ROOT, topic)
    infos, uploads, searches = sha.load_shorts_cache(cache)
    if not infos:
        return None
    channels, ac = sha.load_main_cache(out_all)
    subs_map = {k: v.get("subs") for k, v in channels.items()}
    known = {i for u in uploads.values() for i in u["ids"]}
    cases = sha.load_cases(sdir, infos, subs_map, now, known)
    searched = {i for d in searches for i in (d.get("ids") or [])}
    alltime = sa.load_alltime_ids(cache, dur="short") - searched
    data, success, _group, rows_all, cases = sha.analyze(topic, topic, cases, infos, uploads, searches, channels, ac, rules, rel, themes, now,
                                                         alltime_ids=alltime)
    return {"rules": rules, "data": data, "success": success, "rows": rows_all, "cases": cases, "searched": searched, "alltime": alltime,
            "fetched": {k: v.get("fetched") for k, v in infos.items()}}


def explore_meta(topic):
    """탐색 목록(웹 화면·cloud_run explore)의 검색 기간·추천 까닭·못 한 기간 없는 검색."""
    res = {}
    for fmt, name in ((LONG, "explore.json"), (SHORT, "shorts_explore.json")):
        p = ROOT / topic / "result" / "web" / name
        d = json.loads(p.read_text("utf-8")) if p.is_file() else {}
        res[fmt] = {"weeks": d.get("weeks"), "generated": d.get("generated"), "pending_alltime": len(d.get("pending_alltime") or []),
                    "reason": {r["id"]: r.get("reason") or "" for r in d.get("rows") or []}}
    return res


# ---------------------------------------------------------------- 묶음
def build(long, shorts, pct, meta):
    """두 형식 → 영상 줄 (묶음 표시·분석 상태) + 묶음 요약. v1.1: 분석 결과의 묶음(①#·②V·③T)과 경계를 그대로 읽음 —
    분석에서 꺼진 묶음(TOP_PCT 0)만 여기서 같은 꼴(상위 pct%)로 셈."""
    rows, summary = [], {}
    for fmt, src in ((LONG, long), (SHORT, shorts)):
        if not src:
            summary[fmt] = None
            continue
        R, data = src["rules"], src["data"]
        S = R["success"] if fmt == LONG else R["shorts"]
        is_success = sa.is_success if fmt == LONG else sha.is_success
        is_fmt = (lambda x: x["minutes"] >= R["vpd_top"]["MIN_MINUTES"]) if fmt == LONG else (lambda x: x.get("is_short"))
        vt, tt = data.get("vpd_top") or {}, data.get("views_top") or {}
        max_age = float(vt.get("max_age_days") or 0)
        recent = (lambda x: x.get("age_days") is not None and x["age_days"] <= max_age) if max_age > 0 else (lambda x: True)
        base = [x for x in src["rows"] if not x["exclude"] and x.get("vpd") is not None and is_fmt(x)]
        pool2 = [x for x in base if x["id"] in src["searched"] and recent(x)]                       # ② = 분석 1-3(쇼츠 1-4)번 모집단
        pool3 = [x for x in base if x["id"] in src["searched"] or x["id"] in src["alltime"]]       # ③ = 1-4(쇼츠 1-5)번 모집단
        thr2 = vt.get("threshold") if vt.get("threshold") is not None else orr.top_pct_threshold([x["vpd"] for x in pool2], pct)
        thr3 = tt.get("threshold") if tt.get("threshold") is not None else orr.top_pct_threshold([x["views"] for x in pool3], pct)
        on2, on3 = vt.get("threshold") is not None, tt.get("threshold") is not None
        ids2, ids3 = {x["id"] for x in pool2}, {x["id"] for x in pool3}
        success_no = {c["id"]: c["no"] for c in src["success"]}
        vno = {c["id"]: c["vno"] for c in vt.get("cases") or []}
        tno = {c["id"]: c["tno"] for c in tt.get("cases") or []}
        case_ids = {c["id"] for c in src["cases"]}
        reason = meta[fmt]["reason"]
        url = "https://www.youtube.com/watch?v=" if fmt == LONG else "https://www.youtube.com/shorts/"

        def st(ok, has, num, prefix, on, cap):
            if not ok:
                return ""
            if not on:
                return "분석 묶음 없음 (규칙 꺼짐) · " + ("자료 받음" if has else "자료 안 받음")
            return f"분석 {prefix}{num}" if num is not None else (f"자료 받음 · 채널당 {cap}개 넘어 묶음 밖" if has else "자료 안 받음")
        seen = set()
        for x in src["rows"]:
            if x["id"] in seen:
                continue
            g1 = (not x["exclude"]) and is_success(x, S)
            g2 = x["id"] in ids2 and thr2 is not None and x["vpd"] >= thr2
            g3 = x["id"] in ids3 and thr3 is not None and x["views"] >= thr3
            has = x["id"] in case_ids
            if not (g1 or g2 or g3 or has or x["id"] in ids3):
                continue
            seen.add(x["id"])
            pub, subs = x.get("published"), x.get("subs")
            f = orr.as_utc(src["fetched"].get(x["id"]))
            rows.append({"형식": fmt, "제목": x["title"], "링크": url + x["id"], "채널": x["channel"],
                         "구독자": subs, "조회수": x["views"], "올린 날": pub.date().isoformat() if pub else "",
                         "지난 날수": round(x["age_days"]) if x.get("age_days") is not None else None,
                         "하루 조회수": x["vpd"], "받은 날": f.date().isoformat() if f else "",
                         "구독자당 하루 조회수": round(x["vpd"] / max(subs, 100), 3) if subs is not None and x.get("vpd") is not None else None,
                         "배수(조회수÷구독자)": x.get("ratio"),
                         "길이(분)": round(x["minutes"], 1), "외국": "외국" if x.get("foreign") else "", "주제 묶음": x.get("theme") or "",
                         "어디서 봄": ("검색" if x["id"] in src["searched"] else "기간 없는 검색" if x["id"] in src["alltime"] else
                                     "승인 자료" if has else "채널 평소 영상"),
                         "자료 받음": "○" if has else "",
                         G1: "○" if g1 else "", f"{G1} 분석": st(g1, has, success_no.get(x["id"]), "#", True, S["CASES_PER_CHANNEL"]),
                         G2: "○" if g2 else "", f"{G2} 분석": st(g2, has, vno.get(x["id"]), "V", on2, vt.get("per_channel", 2)),
                         G3: "○" if g3 else "", f"{G3} 분석": st(g3, has, tno.get(x["id"]), "T", on3, tt.get("per_channel", 2)),
                         "탐색 추천 까닭": reason.get(x["id"], ""), "뺀 까닭": x.get("exclude") or "", "_id": x["id"]})
        mine = [r for r in rows if r["형식"] == fmt]
        ages3 = [x["age_days"] for x in pool3 if x.get("age_days") is not None]
        rule1 = (sa.success_rule_text(S).replace("**", "") + (f"·{S['MIN_MINUTES']:g}분 이상" if fmt == LONG else f"·{S['MAX_SECONDS']}초 이하 세로")
                 + f"·채널당 {S['CASES_PER_CHANNEL']}개 (점수 큰 순)"
                 + (" · 무게 1+log2(점수÷기준)로 성공 사례 몫(가중)" if str(S.get("WEIGHT") or "") == "score" else ""))
        summary[fmt] = {
            "weeks": meta[fmt]["weeks"], "generated": meta[fmt]["generated"], "pending_alltime": meta[fmt]["pending_alltime"],
            "pool2": len(pool2), "pool3": len(pool3), "alltime": len(src["alltime"]), "rows": len(src["rows"]), "cases": len(case_ids),
            "oldest3": round(max(ages3)) if ages3 else None, "older_365": sum(1 for a in ages3 if a > 365), "max_age2": max_age,
            "thr_vpd": thr2, "thr_views": thr3, "on2": on2, "on3": on3, "pct": pct, "rule1": rule1,
            "per_day": float(S.get("MIN_VPD_SUBS") or 0) > 0, "weighted": str(S.get("WEIGHT") or "") == "score",
            **{f"{g}_{k}": v for g in (G1, G2, G3) for k, v in (
                ("pass", sum(1 for r in mine if r[g])), ("has", sum(1 for r in mine if r[g] and r["자료 받음"])),
                ("in", sum(1 for r in mine if str(r[f"{g} 분석"]).startswith("분석 ") and not str(r[f"{g} 분석"]).startswith("분석 묶음"))))},
            "g1_uploads": sum(1 for r in mine if r[G1] and r["어디서 봄"] == "채널 평소 영상"),
            "last_fetch": max((r["받은 날"] for r in mine if r["받은 날"]), default=""),
        }
        s = summary[fmt]
        log("GROUP", fmt=fmt, pool2=len(pool2), pool3=len(pool3), alltime=len(src["alltime"]), oldest3=s["oldest3"], older_365=s["older_365"],
            thr_vpd=thr2, thr_views=thr3, g1=f"{s[G1 + '_pass']}/{s[G1 + '_has']}/{s[G1 + '_in']}",
            g2=f"{s[G2 + '_pass']}/{s[G2 + '_has']}/{s[G2 + '_in']}", g3=f"{s[G3 + '_pass']}/{s[G3 + '_has']}/{s[G3 + '_in']}",
            pending_alltime=s["pending_alltime"], note="통과/자료 받음/분석 묶음")
    return rows, summary


def check_rows(summary, long_rules):
    """확인 결과 — 사용자 기준 × 형식 6줄 (v1.1: 분석 결과의 묶음을 그대로 판정)."""
    out = []
    want = {G1: "구독자에 비해 업로드 대비 조회수가 높은 영상", G2: "구독자 상관없이 업로드 대비 조회수가 높은 영상", G3: "업로드 기간 상관없이 조회수가 높은 영상"}
    where = {(G1, LONG): "analysis.md 1번", (G1, SHORT): "shorts/analysis.md 1번", (G2, LONG): "analysis.md 1-3번", (G2, SHORT): "shorts/analysis.md 1-4번",
             (G3, LONG): "analysis.md 1-4번", (G3, SHORT): "shorts/analysis.md 1-5번"}
    for g in (G1, G2, G3):
        for fmt in (LONG, SHORT):
            s = summary.get(fmt)
            if not s:
                out.append({"분석 대상": g, "형식": fmt, "사용자 기준": want[g], "판정": "❌ 자료 없음"})
                continue
            p, h, i = s[f"{g}_pass"], s[f"{g}_has"], s[f"{g}_in"]
            miss = f"기준 통과 {p}개 중 자료를 안 받은 {p - h}개는 분석에 없음 (탐색 추천 — 승인하면 들어감)." if p - h else ""
            if g == G1:
                rule = s["rule1"]
                verdict = "✅ 들어감 (구독자당 하루 조회수)" if s["per_day"] else "⚠️ 들어감 · 날수로 나누지 않음 (배수)"
                gap = ((f"채널 평소 영상(같은 채널 기준선)으로만 본 것 {s['g1_uploads']}개 포함. " if s["g1_uploads"] else "") + miss
                       + (" 하루 조회수 = 받은 때 조회수 ÷ 받은 때까지 날수." if s["per_day"] else ""))
                fix = "" if s["per_day"] else "[success] MIN_VPD_SUBS"
            elif g == G2:
                rule = (f"검색에서 찾은 주제 관련 {'롱폼(4분 이상)' if fmt == LONG else '쇼츠'}"
                        + (f" 중 최근 {s['max_age2']:g}일 안" if s["max_age2"] else "") + f" {s['pool2']:,}개 안 하루 조회수 상위 {s['pct']:g}% "
                        f"(경계 {s['thr_vpd'] or 0:,.0f}회/일) · 채널당 2개")
                verdict = "✅ 들어감" if s["on2"] else "❌ 묶음 꺼짐"
                gap = miss
                fix = "남은 추천 승인 (댓글 받기)" if p - h else ""
            else:
                rule = (f"검색(기간 있는 + 기간 없는 검색) 주제 관련 {'롱폼' if fmt == LONG else '쇼츠'} {s['pool3']:,}개 안 조회수 상위 {s['pct']:g}% "
                        f"(경계 {s['thr_views'] or 0:,.0f}회) · 올린 날짜 상관없이 · 채널당 2개")
                done = s["alltime"] > 0
                verdict = ("✅ 들어감" if done and not s["pending_alltime"] else
                           "⚠️ 들어감 · 기간 없는 검색 일부만" if done else "⚠️ 묶음 있음 · 기간 없는 검색 아직") if s["on3"] else "❌ 묶음 꺼짐"
                gap = (f"기간 없는 검색으로만 찾은 영상 {s['alltime']:,}개 · 모집단에서 가장 오래된 영상 {s['oldest3']}일 · 1년 넘은 {s['older_365']}개"
                       + (f" · 못 한 기간 없는 검색 {s['pending_alltime']}개 (할당량이 2,000단위 넘게 남을 때 이어서)" if s["pending_alltime"] else "")
                       + (". " + miss if miss else "."))
                fix = ("할당량이 남는 날 탐색·이어서 (cloud_run explore/resume)" if s["pending_alltime"] or not done else "") + \
                      ((" · " if s["pending_alltime"] or not done else "") + "남은 추천 승인" if p - h else "")
            out.append({"분석 대상": g, "형식": fmt, "사용자 기준": want[g], "지금 규칙": rule, "기준 통과 (모집단·자료 안)": p, "자료 받음 (승인)": h,
                        "분석 묶음에 듦": i, "분석 결과 위치": where[(g, fmt)], "판정": verdict, "다른 점·빠진 것": gap, "고치려면": fix})
    return out


def pool_rows(summary):
    out = []
    for fmt in (LONG, SHORT):
        s = summary.get(fmt)
        if not s:
            continue
        out.append({"형식": fmt, "검색 기간(주)": s["weeks"], "탐색 목록 만든 때": s["generated"], "② 모집단(최근·검색)": s["pool2"],
                    "③ 모집단(기간 없는 검색 포함)": s["pool3"], "기간 없는 검색으로만 찾은 영상": s["alltime"],
                    "못 한 기간 없는 검색": s["pending_alltime"], "분석에 쓴 영상 전부(채널 평소 영상 포함)": s["rows"], "자료 받은 영상(승인)": s["cases"],
                    "③ 모집단에서 가장 오래된 영상(일)": s["oldest3"], "1년 넘은 영상": s["older_365"],
                    f"② 경계 하루 조회수(상위 {s['pct']:g}%)": s["thr_vpd"], f"③ 경계 조회수(상위 {s['pct']:g}%)": s["thr_views"]})
    return out


def overlap_rows(rows):
    out = []
    for fmt in (LONG, SHORT):
        mine = [r for r in rows if r["형식"] == fmt]
        if not mine:
            continue
        c = Counter((bool(r[G1]), bool(r[G2]), bool(r[G3])) for r in mine)
        for (a, b, d), n in sorted(c.items(), key=lambda kv: -kv[1]):
            if not (a or b or d):
                continue
            out.append({"형식": fmt, G1: "○" if a else "", G2: "○" if b else "", G3: "○" if d else "", "영상 수": n})
    return out


# ---------------------------------------------------------------- 엑셀
FMT_INT = ("구독자", "조회수", "하루 조회수", "지난 날수", "기준 통과 (모집단·자료 안)", "자료 받음 (승인)", "분석 묶음에 듦", "영상 수",
           "1년 넘은 영상", "자료 받은 영상(승인)", "② 모집단(최근·검색)", "③ 모집단(기간 없는 검색 포함)", "기간 없는 검색으로만 찾은 영상",
           "못 한 기간 없는 검색", "분석에 쓴 영상 전부(채널 평소 영상 포함)", "③ 모집단에서 가장 오래된 영상(일)")


def write_xlsx(path, sheets, guide):
    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame({"안내": guide}).to_excel(xw, sheet_name="안내", index=False)
        for name, rows in sheets:
            pd.DataFrame(rows).to_excel(xw, sheet_name=name, index=False)
        wb = xw.book
        fill, wrap = PatternFill("solid", fgColor="DDEBF7"), Alignment(wrap_text=True, vertical="top")
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            if ws.max_row > 1 and ws.title != "안내":
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font, cell.fill, cell.alignment = Font(bold=True), fill, wrap
            for col in range(1, ws.max_column + 1):
                name = str(ws.cell(row=1, column=col).value or "")
                wide = name in ("안내", "제목", "지금 규칙", "다른 점·빠진 것", "고치려면", "사용자 기준", "분석 결과 위치") or name.endswith(" 분석")
                ws.column_dimensions[get_column_letter(col)].width = (70 if name in ("안내", "다른 점·빠진 것", "지금 규칙") else 45 if wide else
                                                                      max(9, min(22, len(name) * 2 + 2)))
                fmt = ("#,##0" if name in FMT_INT or name.startswith("② 경계") or name.startswith("③ 경계") else
                       "0.00" if name == "배수(조회수÷구독자)" else "0.000" if name == "구독자당 하루 조회수" else None)
                if not (fmt or wide or name == "링크"):
                    continue
                for (c,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                    if fmt:
                        c.number_format = fmt
                    if wide:
                        c.alignment = wrap
                    if name == "링크" and c.value:
                        c.hyperlink, c.font = c.value, Font(color="0563C1", underline="single")
    log("XLSX", file=str(path), sheets=len(sheets) + 1, kb=round(path.stat().st_size / 1024))


def write_md(path, checks, pools, now):
    L = [f"# 분석 대상 세 묶음 확인 (`tools/target_check.py` {VERSION} · {now.date().isoformat()}) — 자세한 영상 목록은 `target_check.xlsx`", "",
         "※ 연구·교육용 — 투자 권유가 아님. 숫자 = 기준 통과 / 자료 받음(승인) / 분석 묶음에 듦.", "",
         "| 분석 대상 | 형식 | 판정 | 통과 / 자료 / 분석 | 다른 점·빠진 것 |", "|---|---|---|---|---|"]
    for r in checks:
        L.append(f"| {r['분석 대상']} | {r['형식']} | {r['판정']} | {r.get('기준 통과 (모집단·자료 안)', '-')} / {r.get('자료 받음 (승인)', '-')} / "
                 f"{r.get('분석 묶음에 듦', '-')} | {str(r.get('다른 점·빠진 것', '')).replace('|', '∣')} |")
    L += ["", "| 형식 | 검색 기간 | ② 모집단 | ③ 모집단 (기간 없는 검색으로만) | ③ 가장 오래된 영상 | 못 한 기간 없는 검색 | ② 경계 하루 조회수 | ③ 경계 조회수 |",
          "|---|---|---|---|---|---|---|---|"]
    for p in pools:
        k2 = next(k for k in p if k.startswith("② 경계"))
        k3 = next(k for k in p if k.startswith("③ 경계"))
        L.append(f"| {p['형식']} | {p['검색 기간(주)']}주 | {p['② 모집단(최근·검색)']:,} | {p['③ 모집단(기간 없는 검색 포함)']:,} ({p['기간 없는 검색으로만 찾은 영상']:,}) | "
                 f"{p['③ 모집단에서 가장 오래된 영상(일)']}일 | {p['못 한 기간 없는 검색']} | {p[k2] or 0:,.0f} | {p[k3] or 0:,.0f} |")
    path.write_text("\n".join(L) + "\n", "utf-8")


def run(topic, pct=None, now=None, out_path=None):
    now = now or datetime.now(timezone.utc)
    t0 = time.time()
    long = load_long(topic, now)
    shorts = load_shorts(topic, now)
    pct = pct if pct is not None else long["rules"]["vpd_top"]["TOP_PCT"]
    meta = explore_meta(topic)
    log("LOAD", topic=topic, long_rows=len(long["rows"]), long_cases=len(long["cases"]), shorts_rows=len(shorts["rows"]) if shorts else 0,
        shorts_cases=len(shorts["cases"]) if shorts else 0, long_weeks=meta[LONG]["weeks"], shorts_weeks=meta[SHORT]["weeks"], pct=pct,
        seconds=round(time.time() - t0, 1))
    rows, summary = build(long, shorts, pct, meta)
    checks, pools = check_rows(summary, long["rules"]), pool_rows(summary)
    order = {LONG: 0, SHORT: 1}
    def pick(col, key):
        return [{k: v for k, v in r.items() if k != "_id"} for r in sorted((r for r in rows if r[col]), key=lambda r: (order[r["형식"]], -(r[key] or 0)))]
    sheets = [("확인 결과", checks), ("모집단", pools),
              ("① 성공 사례", pick(G1, "배수(조회수÷구독자)")), ("② 최근 트렌드", pick(G2, "하루 조회수")), ("③ 조회수 상위", pick(G3, "조회수")),
              ("겹침", overlap_rows(rows)),
              ("전체 목록", [{k: v for k, v in r.items() if k != "_id"} for r in sorted(rows, key=lambda r: (order[r["형식"]], -(r["하루 조회수"] or 0)))])]
    out = ROOT / topic / "result" / "outliers"
    out_path = Path(out_path) if out_path else out / "target_check.xlsx"
    guide = [
        f"자동 생성 — tools/target_check.py {VERSION} ({now.isoformat(timespec='minutes')}). 연구·교육용 — 투자 권유가 아님. 할당량 0 (받아 둔 자료만).",
        "사용자 지시 (2026-10-28): 분석 대상(롱폼·쇼츠 종합) ① 구독자에 비해 업로드 대비 조회수가 높은 영상(성공 사례) ② 구독자 상관없이 업로드 대비 조회수가 높은 영상"
        "(최근 트렌드) ③ 업로드 기간 상관없이 조회수가 높은 영상(최근 트렌드에 비슷한 과거 사례) — 이렇게 분석하고 있는지 확인.",
        "셈하는 법: 분석 코드(success_analysis·shorts_analysis)를 그대로 불러 같은 모집단·규칙·묶음으로 셈 — 분석 결과 파일은 바꾸지 않음 (오늘 날짜로 다시 셈).",
        "① = 성공 사례 규칙 그대로 (확인 결과 '지금 규칙' — 2026-10-28 부터 구독자당 하루 조회수 ≥ 0.1). ② = 하루 조회수 상위 N% (롱폼 1-3번·쇼츠 1-4번 — 최근 26주 안). "
        "③ = 조회수 상위 N% (롱폼 1-4번·쇼츠 1-5번 — 기간 없는 검색 포함, 올린 날짜 상관없이).",
        "모집단 = 탐색 검색에서 찾은 주제 관련 영상 중 뺄 까닭(라이브·주제 밖·광고 의심·쇼츠 부풀림 의심 등)이 없는 것. 롱폼은 4분 이상, 쇼츠는 쇼츠인 것. "
        "기간 없는 검색(전체 기간 조회수 순)은 할당량이 2,000단위 넘게 남는 날에만 하므로 처음 며칠은 ③ 의 과거 영상이 적다 ('모집단' 시트의 못 한 기간 없는 검색).",
        "'○' = 기준 통과. '… 분석' 칸: '분석 #n' = 성공 사례 n번 · '분석 V n' = 하루 조회수 상위 · '분석 T n' = 조회수 상위 · '자료 받음 · 채널당 2개 넘어 묶음 밖' · "
        "'자료 안 받음' (승인하지 않아 설명·댓글이 없음 — 탐색 추천을 승인하면 들어감).",
        "하루 조회수 = 받은 때 조회수 ÷ 올린 날부터 받은 때까지 날수 (2026-10-28 부터 — 예전은 ÷ 오늘까지라 자료가 오래되면 낮게 셈). 구독자당 하루 조회수 = 하루 조회수 ÷ 구독자"
        "(100 미만은 100). 어디서 봄 = 검색 · 기간 없는 검색 · 승인 자료 · 채널 평소 영상(같은 채널 기준선으로 받은 최근 영상).",
        "탐색 추천 까닭 = 탐색 목록이 체크해 둔 까닭 (성공 사례 · 하루 조회수 상위 · 조회수 상위 · 성공 쇼츠).",
        "한계: 조회수·구독자는 받은 때 값. 검색은 YouTube 검색 결과(조회수 순·검색어마다 최대 50개)라 모든 영상이 아님.",
        f"다시 만들기: python3 tools/target_check.py {topic}" + (f" --pct {pct:g}" if pct != long["rules"]["vpd_top"]["TOP_PCT"] else ""),
    ]
    write_xlsx(out_path, sheets, guide)
    write_md(out_path.with_suffix(".md"), checks, pools, now)
    log("DONE", seconds=round(time.time() - t0, 1), xlsx=str(out_path), rows=len(rows),
        verdict=" | ".join(f"{r['분석 대상'][:1]}{r['형식']}={r['판정']}" for r in checks))
    return {"xlsx": str(out_path), "checks": checks, "pools": pools, "rows": len(rows)}


def main(argv=None):
    ap = argparse.ArgumentParser(description="분석 대상 세 묶음(성공 사례·최근 트렌드·조회수 상위) 확인 엑셀 — 롱폼·쇼츠 종합 (할당량 0)")
    ap.add_argument("topic")
    ap.add_argument("--pct", type=float, default=None, help="②·③ 상위 몇 %% (기본 = analysis_rules.toml [vpd_top] TOP_PCT)")
    ap.add_argument("--out", default="", help="엑셀 경로 (기본 <주제>/result/outliers/target_check.xlsx)")
    ap.add_argument("--log-level", default="INFO", choices=["DEBUG", "INFO", "WARNING"])
    a = ap.parse_args(argv)
    logging.basicConfig(level=getattr(logging, a.log_level), format="%(message)s", stream=sys.stderr)
    log("START", topic=a.topic, pct=a.pct if a.pct is not None else "기본", tool=f"target_check.py {VERSION}")
    run(a.topic, pct=a.pct, out_path=a.out or None)


if __name__ == "__main__":
    main()
